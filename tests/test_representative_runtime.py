from __future__ import annotations

from contextlib import redirect_stdout
import io
import os
from pathlib import Path
import subprocess
import sys
import tempfile
import time
import unittest
from unittest.mock import patch

from openpyxl import load_workbook

import uruchom_klientow_reprezentatywnych as launcher
from representative_clients.portfolio_cli import main
from representative_clients.runtime import (
    RunBudget, RunTimeLimitError, run_portfolio, supervise, validate_run_seconds,
)
from tests.test_representative_launcher import make_portfolio


class RuntimeTests(unittest.TestCase):
    def test_invalid_or_over_24h_limits_are_rejected_before_spawn(self):
        for seconds in (0, -1, 86401, float("inf"), float("nan")):
            with self.subTest(seconds=seconds):
                with self.assertRaises(ValueError):
                    validate_run_seconds(seconds)
                with patch("representative_clients.runtime.subprocess.Popen") as spawn:
                    code = run_portfolio(
                        ["--input", "unused.xlsx", "--consents", "unused.xlsx", "--total-time-limit", str(seconds)],
                        on_line=lambda line: None,
                    )
                    self.assertEqual(code, 2)
                    spawn.assert_not_called()

    def test_remaining_budget_is_shared_and_preserves_export_margin(self):
        with patch("representative_clients.runtime.time.monotonic", return_value=100) as clock:
            budget = RunBudget(100)
            self.assertEqual(budget.group_seconds(30, 4), 22.5)
            clock.return_value = 150
            self.assertEqual(budget.group_seconds(30, 2), 20)
            self.assertEqual(budget.group_seconds(5, 2), 5)
            clock.return_value = 190
            with self.assertRaises(RunTimeLimitError):
                budget.group_seconds(30, 1)
            clock.return_value = 200
            with self.assertRaises(RunTimeLimitError):
                budget.check()

    def test_silent_worker_is_killed_and_reaped_at_deadline(self):
        processes = []
        real_popen = subprocess.Popen

        def spawn(*args, **kwargs):
            child = real_popen(*args, **kwargs)
            processes.append(child)
            return child

        lines = []
        started = time.monotonic()
        with patch("representative_clients.runtime.subprocess.Popen", side_effect=spawn):
            code = supervise([sys.executable, "-c", "import time; time.sleep(30)"], 0.3, lines.append)
        self.assertEqual(code, 124)
        self.assertLess(time.monotonic() - started, 5)
        self.assertIsNotNone(processes[0].poll())
        self.assertIn("STOP", "".join(lines))

    def test_worker_can_finish_early_and_output_and_exit_code_are_preserved(self):
        lines = []
        started = time.monotonic()
        code = supervise([sys.executable, "-c", "print('gotowe'); raise SystemExit(7)"], 86400, lines.append)
        self.assertEqual(code, 7)
        self.assertEqual("".join(lines), "gotowe\n")
        self.assertLess(time.monotonic() - started, 5)

    def test_continuous_output_does_not_defeat_deadline(self):
        code = supervise(
            [sys.executable, "-u", "-c", "import time\nwhile True:\n print('postep'); time.sleep(0.001)"],
            0.3, lambda line: None,
        )
        self.assertEqual(code, 124)

    def test_keyboard_interrupt_stops_worker(self):
        def interrupted(line):
            if line.startswith("ready"):
                raise KeyboardInterrupt
        code = supervise(
            [sys.executable, "-u", "-c", "import time; print('ready'); time.sleep(30)"],
            10, interrupted,
        )
        self.assertEqual(code, 130)

    def test_cooperative_deadline_stops_before_reading_inputs(self):
        with patch("representative_clients.consent.load_consent_workbook") as load, redirect_stdout(io.StringIO()):
            code = main(
                ["--input", "unused.xlsx", "--consents", "unused.xlsx", "--validate-full"],
                deadline=time.monotonic() - 1,
            )
        self.assertEqual(code, 124)
        load.assert_not_called()

    def test_checkpoint_survives_timeout_in_next_branch(self):
        from representative_clients.portfolio_cli import load_gams_workbook

        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            args = make_portfolio(root)
            output = root / "wynik.xlsx"

            def read(*args, **kwargs):
                if Path(args[0]).name == "model_lublin.xlsx":
                    raise RunTimeLimitError("Test limitu w następnym OSD")
                return load_gams_workbook(*args, **kwargs)

            with patch("representative_clients.portfolio_cli.load_gams_workbook", side_effect=read), redirect_stdout(io.StringIO()):
                code = main([*args, "--output", str(output), "--time-limit", "10"])
            self.assertEqual(code, 124)
            self.assertFalse(output.exists())
            checkpoints = list(root.glob("wynik_czesci_*/*.xlsx"))
            self.assertEqual(len(checkpoints), 1)
            book = load_workbook(checkpoints[0], read_only=True, data_only=True)
            try:
                self.assertEqual(book["Podsumowanie"].max_row, 2)
                self.assertGreater(book["Zastepcy"].max_row, 1)
                self.assertIn("CZĘŚCIOWY", book["Tekst_do_maila"]["A2"].value)
                config = dict(book["Konfiguracja"].iter_rows(min_row=2, values_only=True))
                self.assertEqual(config["limit_calego_przebiegu_s"], 86400)
                self.assertIn("CZĘŚCIOWY", config["zakres_raportu"])
            finally:
                book.close()

    def test_launcher_runs_supervised_pipeline_and_saves_all_checkpoints(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            args = make_portfolio(root)
            output = root / "wynik.xlsx"
            log = root / "przebieg.log"
            with redirect_stdout(io.StringIO()):
                code = launcher.execute_run(
                    [*args, "--output", str(output), "--time-limit", "10", "--total-time-limit", "60"], log,
                )
            self.assertEqual(code, 0)
            self.assertTrue(output.exists())
            self.assertEqual(len(list(root.glob("wynik_czesci_*/*.xlsx"))), 8)
            self.assertIn("Limit całego przebiegu", log.read_text(encoding="utf-8"))

    def test_cli_also_enforces_deadline_during_worker_startup(self):
        result = subprocess.run(
            [sys.executable, "-m", "representative_clients.portfolio_cli", "--input", "unused.xlsx",
             "--consents", "unused.xlsx", "--validate-full", "--total-time-limit", "0.001"],
            capture_output=True, text=True, timeout=10,
        )
        self.assertEqual(result.returncode, 124)
        self.assertIn("STOP", result.stdout)

    def test_relative_inputs_keep_callers_working_directory(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            args = make_portfolio(root)
            args = [Path(item).name if item.endswith(".xlsx") else item for item in args]
            previous = Path.cwd()
            lines = []
            try:
                os.chdir(root)
                code = run_portfolio([*args, "--validate-only"], on_line=lines.append)
            finally:
                os.chdir(previous)
            self.assertEqual(code, 0, "".join(lines))
            self.assertIn("Kontrola list zakończona", "".join(lines))


if __name__ == "__main__":
    unittest.main()
