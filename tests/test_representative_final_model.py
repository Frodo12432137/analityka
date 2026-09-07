from __future__ import annotations

import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

import numpy as np
from openpyxl import Workbook, load_workbook

from representative_clients.consent import canonical_branch, load_consent_workbook
from representative_clients.model import GroupResult, SelectionConfig, SelectionData, SelectionResult
from representative_clients.portfolio_data import (
    allocate_branch_targets,
    allocate_group_quotas,
    merge_selection_data,
)
from representative_clients.portfolio_report import BranchOutput, write_portfolio_workbook
from representative_clients.replacement import choose_unique_reserves, optimal_nonnegative_weight
from representative_clients.solver import solve_selection


class ConsentTests(unittest.TestCase):
    def test_reads_real_header_names_and_maps_osd(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "zgody.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "page"
            sheet.append(["Kod PP Aktualny", "OSD Nazwa - PP Aktualny"])
            sheet.append(["590000000000000001", "PGE Dystrybucja S.A. Oddział Białystok"])
            sheet.append(["590000000000000002", "PGED_Lodz_Miasto"])
            workbook.save(path)
            index = load_consent_workbook(path)
        self.assertEqual(index.unique_count, 2)
        self.assertIn("590000000000000001", index.by_branch["bialystok"])
        self.assertIn("590000000000000002", index.by_branch["lodz_miasto"])

    def test_all_eight_branch_names_are_recognized(self) -> None:
        values = {
            "PGED_Bialystok": "bialystok",
            "Oddział Lublin": "lublin",
            "Łódź Miasto": "lodz_miasto",
            "Lodz Teren": "lodz_teren",
            "Rzeszów": "rzeszow",
            "Skarżysko-Kamienna": "skarzysko_kamienna",
            "Warszawa": "warszawa",
            "Zamość": "zamosc",
        }
        for value, expected in values.items():
            self.assertEqual(canonical_branch(value), expected)


class AllocationTests(unittest.TestCase):
    def test_branch_allocation_is_exact_and_reserves_capacity_is_kept(self) -> None:
        result = allocate_branch_targets({"bialystok": 10, "lublin": 20}, 9)
        self.assertEqual(sum(result.values()), 9)
        self.assertLessEqual(result["bialystok"], 5)
        self.assertLessEqual(result["lublin"], 10)

    def test_group_quota_respects_required_clients(self) -> None:
        data = SelectionData(
            branch="PGED_Bialystok",
            client_ids=("A", "B", "C", "D", "E", "F"),
            periods=(("2026-01-01", 1),),
            client_balance=np.ones((1, 6)),
            branch_balance=np.ones(1),
            required_weights={"A": 1.0, "C": 1.0},
        )
        quotas = allocate_group_quotas(data, 3, 2)
        self.assertEqual(sum(quotas), 3)
        self.assertGreaterEqual(quotas[0], 2)

    def test_merge_deduplicates_identical_client(self) -> None:
        data = SelectionData(
            branch="PGED_Zamosc",
            client_ids=("A",),
            periods=(("2026-01-01", 1),),
            client_balance=np.asarray([[2.0]]),
            branch_balance=np.asarray([4.0]),
        )
        merged = merge_selection_data((data, data))
        self.assertEqual(merged.client_ids, ("A",))
        np.testing.assert_allclose(merged.client_balance, [[2.0]])


class FinalSolverTests(unittest.TestCase):
    def _data(self) -> SelectionData:
        return SelectionData(
            branch="PGED_Bialystok",
            client_ids=("A", "B", "C", "D", "E", "F"),
            periods=(("2026-01-01", 1), ("2026-01-01", 2)),
            client_balance=np.asarray(
                [[1.0, 0.8, 0.5, 0.4, 1.2, 0.7], [2.0, 1.5, 1.0, 0.9, 2.5, 1.4]]
            ),
            branch_balance=np.asarray([3.0, 6.0]),
        )

    def test_exact_group_count_is_enforced(self) -> None:
        result = solve_selection(
            self._data(),
            SelectionConfig(
                max_clients_per_group=2,
                group_count=1,
                exact_clients_per_group=(2,),
                max_contribution_share=10.0,
                time_limit_seconds=30,
                mip_relative_gap=0,
            ),
        )
        self.assertEqual(result.selected_count, 2)

    def test_remaining_time_limit_is_passed_to_native_solver(self) -> None:
        from scipy.optimize import milp

        with patch("scipy.optimize.milp", wraps=milp) as native:
            result = solve_selection(
                self._data(),
                SelectionConfig(
                    max_clients_per_group=2, group_count=1,
                    exact_clients_per_group=(2,), max_contribution_share=10,
                    time_limit_seconds=30, mip_relative_gap=0,
                ),
                time_limit_provider=lambda: 2.0,
            )
        self.assertEqual(result.selected_count, 2)
        self.assertEqual(native.call_args.kwargs["options"]["time_limit"], 2.0)

    def test_weighted_median_weight_is_exact_for_scaled_profile(self) -> None:
        weight, error = optimal_nonnegative_weight(
            np.asarray([1.0, 2.0, -1.0]), np.asarray([3.0, 6.0, -3.0])
        )
        self.assertAlmostEqual(weight, 3.0)
        self.assertAlmostEqual(error, 0.0)

    def test_reserves_are_unique_and_report_is_auditable(self) -> None:
        data = self._data()
        group = GroupResult(
            group="w1",
            candidate_ids=data.client_ids,
            selected_weights={"A": 2.0, "B": 1.0},
            binary_values=np.asarray([1, 1, 0, 0, 0, 0], dtype=float),
            weights=np.asarray([2, 1, 0, 0, 0, 0], dtype=float),
            model_balance=data.client_balance[:, 0] * 2 + data.client_balance[:, 1],
            residual=np.asarray([-0.2, -0.5]),
            objective_value=0.7,
            status=0,
            status_text="Optimal",
            has_solution=True,
            optimal=True,
            solve_seconds=0.1,
        )
        selection = SelectionResult(
            branch=data.branch,
            periods=data.periods,
            branch_balance=data.branch_balance,
            config=SelectionConfig(max_clients_per_group=2, group_count=1),
            groups=(group,),
            required_weights={},
            excluded=frozenset(),
        )
        replacements = choose_unique_reserves(data, selection, initial_shortlist=1)
        self.assertEqual(len(replacements.matches), 2)
        self.assertEqual(len({item.reserve_id for item in replacements.matches}), 2)
        self.assertFalse({item.reserve_id for item in replacements.matches} & {"A", "B"})

        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "final.xlsx"
            write_portfolio_workbook(
                (
                    BranchOutput(
                        selection=selection,
                        replacements=replacements,
                        source_files=("model.xlsx",),
                        source_clients=6,
                        consent_clients=6,
                        legacy_excluded=0,
                        eligible_clients=6,
                        target_clients=2,
                    ),
                ),
                path,
                2,
            )
            workbook = load_workbook(path, read_only=True, data_only=True)
            try:
                self.assertIn("Tekst_do_maila", workbook.sheetnames)
                self.assertEqual(workbook["Podsumowanie"]["C2"].value, 2)
                self.assertEqual(workbook["Zastepcy"].max_row, 3)
            finally:
                workbook.close()


if __name__ == "__main__":
    unittest.main()
