from __future__ import annotations

from contextlib import redirect_stdout, redirect_stderr
import io
import os
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch

from openpyxl import Workbook, load_workbook

import uruchom_klientow_reprezentatywnych as launcher
from representative_clients.consent import BRANCH_ORDER, technical_branch
from representative_clients.portfolio_cli import main


def make_portfolio(root: Path, total: int = 500, eligible_per_branch: int = 128) -> list[str]:
    """8 OSD, dopuszczone PPE + 1 bez zgody na OSD; pełny schemat v8."""
    args = []
    consent_book = Workbook()
    consent_book.active.append(["Kod PP Aktualny", "OSD Nazwa - PP Aktualny"])
    for branch_index, key in enumerate(BRANCH_ORDER):
        name = technical_branch(key)
        book = Workbook()
        book.remove(book.active)
        target = book.create_sheet("Oddzial")
        target.append(["Bilans"])
        target.append(["Oddzial", "Data", 1, 2])
        target.append([name, "2025-01-13", 100, -50])
        clients = book.create_sheet("Klienci")
        clients.append(["PPE"])
        profiles = book.create_sheet("Klienci_part_1")
        profiles.append(["Profile"])
        profiles.append(["Oddzial", "Klient", "Data", 1, 2])
        other = book.create_sheet("Klienci_part_2")
        other.append(["Profile"])
        other.append(["Oddzial", "Klient", "Data", 1, 2])
        restrictions = book.create_sheet("Ograniczenia")
        restrictions.append(["Wykluczeni", None, None, None, "Wymagani", "Waga"])
        for index in range(eligible_per_branch + 1):
            code = f"TEST{branch_index:02d}{index:05d}"
            clients.append([code])
            profiles.append([name, code, "2025-01-13", 2, -1])
            if index < eligible_per_branch:
                consent_book.active.append([code, name])
        path = root / f"model_{key}.xlsx"
        book.save(path)
        book.close()
        args.extend(("--input", str(path)))
    consent_path = root / "zgody.xlsx"
    consent_book.save(consent_path)
    consent_book.close()
    args.extend(("--consents", str(consent_path), "--total-clients", str(total)))
    return args


class LauncherTests(unittest.TestCase):
    def test_argument_error_does_not_claim_a_workbook_was_written(self):
        with tempfile.TemporaryDirectory() as directory:
            stdout = io.StringIO()
            with redirect_stdout(stdout):
                code = launcher.execute_run([], Path(directory) / "error.log")
            self.assertEqual(code, 2)
            self.assertIn("Błąd (kod 2)", stdout.getvalue())
            self.assertNotIn("Wynik zapisany", stdout.getvalue())

    def test_child_runs_from_other_working_directory_and_keeps_log(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            log = root / "help.log"
            previous = Path.cwd()
            try:
                os.chdir(root)
                with redirect_stdout(io.StringIO()):
                    code = launcher.execute_run(["--help"], log)
            finally:
                os.chdir(previous)
            self.assertEqual(code, 0)
            self.assertIn("--validate-full", log.read_text(encoding="utf-8"))

    def test_missing_dependency_has_install_command_and_does_not_open_form(self):
        stderr = io.StringIO()
        with patch.object(launcher.importlib, "import_module", side_effect=ImportError("scipy")), \
             patch.object(launcher, "choose_run") as form, redirect_stderr(stderr):
            code = launcher.run()
        self.assertEqual(code, 1)
        self.assertIn("-m pip install -r", stderr.getvalue())
        self.assertIn("requirements.txt", stderr.getvalue())
        form.assert_not_called()


class PortfolioExecutionTests(unittest.TestCase):
    def test_part_with_no_consented_clients_does_not_block_other_parts(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            args = make_portfolio(root)
            extra = root / "bialystok_dodatkowa_czesc.xlsx"
            book = load_workbook(root / "model_bialystok.xlsx")
            book["Klienci"].delete_rows(2, 129)
            book["Klienci"].append(["BRAK_ZGODY_W_TEJ_CZESCI"])
            book["Klienci_part_1"].delete_rows(3, 129)
            book["Klienci_part_1"].append(["PGED_Bialystok", "BRAK_ZGODY_W_TEJ_CZESCI", "2025-01-13", 2, -1])
            book.save(extra)
            book.close()
            with redirect_stdout(io.StringIO()) as stdout:
                code = main([*args, "--input", str(extra), "--validate-full"])
            self.assertEqual(code, 0)
            self.assertIn("pominięto profile tej części", stdout.getvalue())

    def test_full_validation_reads_last_profile_and_never_solves(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            args = make_portfolio(root)
            with patch("representative_clients.portfolio_cli.solve_selection") as solve, \
                 patch("representative_clients.portfolio_cli.write_portfolio_workbook") as write, \
                 redirect_stdout(io.StringIO()):
                self.assertEqual(main([*args, "--validate-full"]), 0)
                solve.assert_not_called()
                write.assert_not_called()
            last_source = root / "model_zamosc.xlsx"
            book = load_workbook(last_source)
            book["Klienci_part_1"]["D3"] = "=1+1"  # brak cache; pełna kontrola musi wykryć
            book.save(last_source)
            book.close()
            with patch("representative_clients.portfolio_cli.solve_selection") as solve, \
                 redirect_stdout(io.StringIO()), redirect_stderr(io.StringIO()) as errors:
                with self.assertRaises(SystemExit) as caught:
                    main([*args, "--validate-full"])
                self.assertEqual(caught.exception.code, 1)
                self.assertIn("formuły bez zapisanej wartości", errors.getvalue())
                solve.assert_not_called()

    def test_eight_branches_produce_exact_distinct_pairs_at_both_count_limits(self):
        for total, eligible in ((500, 128), (1100, 280)):
            with self.subTest(total=total), tempfile.TemporaryDirectory() as directory:
                root = Path(directory)
                args = make_portfolio(root, total, eligible)
                if total == 1100:
                    # Brak --total-clients sprawdza domyślny pełny przebieg.
                    args = args[:-2]
                output = root / "wynik.xlsx"
                stdout = io.StringIO()
                with redirect_stdout(stdout):
                    code = main([*args, "--output", str(output), "--time-limit", "10", "--mip-gap", "0"])
                self.assertEqual(code, 0)
                self.assertIn("Budżet solvera:", stdout.getvalue())
                self.assertIn("Cały etap OSD", stdout.getvalue())
                self.assertIn("w1:", stdout.getvalue())
                book = load_workbook(output, read_only=True, data_only=True)
                try:
                    primary = {row[2] for row in book["Wybrani"].iter_rows(min_row=2, values_only=True)}
                    pairs = list(book["Zastepcy"].iter_rows(min_row=2, values_only=True))
                    reserves = {row[4] for row in pairs}
                    self.assertEqual(book["Wybrani"].max_row, total + 1)
                    self.assertEqual(len(primary), total)
                    self.assertEqual(len(pairs), total)
                    self.assertEqual({row[2] for row in pairs}, primary)
                    self.assertEqual(len(reserves), total)
                    self.assertFalse(primary & reserves)
                    self.assertFalse(any(code.endswith(f"{eligible:05d}") for code in primary | reserves))
                    self.assertTrue(all(row[2][4:6] == row[4][4:6] for row in pairs))
                    self.assertEqual(book["Podsumowanie"].max_row, 9)
                finally:
                    book.close()

    def test_quick_check_rejects_required_ppe_without_consent(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            args = make_portfolio(root)
            source = root / "model_bialystok.xlsx"
            book = load_workbook(source)
            book["Ograniczenia"].append([None, None, None, None, "TEST0000128", 1])
            book.save(source)
            book.close()
            with patch("representative_clients.portfolio_cli.solve_selection") as solve, \
                 redirect_stdout(io.StringIO()), redirect_stderr(io.StringIO()) as errors:
                with self.assertRaises(SystemExit):
                    main([*args, "--validate-only"])
                self.assertIn("wymaganych PPE", errors.getvalue())
                solve.assert_not_called()


if __name__ == "__main__":
    unittest.main()
