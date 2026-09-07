"""Sztuczne dane układu z profilami w Klienci; bez danych ze zdjęć użytkownika."""
from contextlib import redirect_stdout
import io
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch

import numpy as np
from openpyxl import Workbook, load_workbook

from representative_clients.legacy_excel import load_gams_workbook
from representative_clients.portfolio_cli import main
from representative_clients.portfolio_data import inspect_workbook_clients
from tests.test_representative_launcher import make_portfolio


def inline_workbook():
    book = Workbook()
    sheet = book.active
    sheet.title = "Klienci"
    sheet.append(["Suma kontrolna — nie dane", None, None, 9999, 8888])
    sheet.append(["Oddział", "Klient", "dataCET", 1, 2])
    for date, a, b in (("2025-01-13", (1, -2), (3, 4)), ("2025-02-09", (2, -3), (5, 6))):
        sheet.append(["PGED_Bialystok", "TEST_A", date, *a])
        sheet.append(["PGED_Bialystok", "TEST_B", date, *b])
    target = book.create_sheet("Oddzial")
    target.append(["Suma kontrolna", None, 7777, 6666])
    target.append(["Oddzial", "dataCET", 1, 2])
    target.append(["PGED_Bialystok", "2025-01-13", 10, -20])
    target.append(["PGED_Bialystok", "2025-02-09", 20, -30])
    limits = book.create_sheet("Ograniczenia")
    limits.append(["Wykluczeni", "Opis", None, None, "Wymagani", "Waga"])
    limits.append(["TEST_B", "Test", None, None, "TEST_A", 1])
    return book


class InlineExcelTests(unittest.TestCase):
    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory()
        self.addCleanup(self.temporary.cleanup)
        self.path = Path(self.temporary.name) / "bialystok_test.xlsx"

    def save(self, book):
        book.save(self.path)
        book.close()

    def test_inline_profiles_deduplicate_ids_by_first_occurrence_not_dates(self):
        self.save(inline_workbook())
        data = load_gams_workbook(self.path)
        self.assertEqual(data.client_ids, ("TEST_A", "TEST_B"))
        self.assertEqual(data.periods, (("2025-01-13", 1), ("2025-01-13", 2), ("2025-02-09", 1), ("2025-02-09", 2)))
        np.testing.assert_array_equal(data.client_balance, [[1, 3], [-2, 4], [2, 5], [-3, 6]])
        np.testing.assert_array_equal(data.branch_balance, [10, -20, 20, -30])
        self.assertEqual(data.excluded, {"TEST_B"})
        self.assertEqual(data.required_weights, {"TEST_A": 1})

    def test_consent_filter_and_quick_check_read_ppe_from_b(self):
        self.save(inline_workbook())
        all_ids, allowed, excluded, required = inspect_workbook_clients(self.path, {"TEST_A"})
        self.assertEqual(all_ids, {"TEST_A", "TEST_B"})
        self.assertEqual(allowed, {"TEST_A"})
        self.assertEqual(excluded, set())
        self.assertEqual(required, {"TEST_A"})
        data = load_gams_workbook(self.path, allowed_client_ids={"TEST_A"})
        self.assertEqual(data.client_ids, ("TEST_A",))
        np.testing.assert_array_equal(data.client_balance[:, 0], [1, -2, 2, -3])

    def test_split_and_inline_layouts_produce_identical_data(self):
        inline = inline_workbook()
        self.save(inline)
        expected = load_gams_workbook(self.path)
        split = inline_workbook()
        split["Klienci"].title = "Klienci_part_1"
        domain = split.create_sheet("Klienci")
        domain.append(["PPE"])
        domain.append(["TEST_A"])
        domain.append(["TEST_B"])
        other = split.create_sheet("Klienci_part_2")
        other.append(["Profile"])
        other.append(["Oddzial", "Klient", "Data", 1, 2])
        self.save(split)
        actual = load_gams_workbook(self.path)
        self.assertEqual(actual.client_ids, expected.client_ids)
        np.testing.assert_array_equal(actual.client_balance, expected.client_balance)

    def test_inline_uncached_formula_in_hour_or_ppe_is_not_silently_empty(self):
        for cell in ("D3", "B3"):
            with self.subTest(cell=cell):
                book = inline_workbook()
                book["Klienci"][cell] = "=1+1"
                self.save(book)
                with self.assertRaisesRegex(ValueError, f"bialystok_test.xlsx: .*formuły.*Klienci!{cell}"):
                    load_gams_workbook(self.path)

    def test_bad_layouts_and_values_include_the_filename(self):
        mutations = (
            (lambda b: b["Klienci"].append(["PGED_Bialystok", "TEST_A", "2025-01-13", 1, -2]), "duplikat"),
            (lambda b: b["Klienci"].__setitem__("B3", 590000000000000001), "jako tekst"),
            (lambda b: b["Klienci"].__setitem__("C2", "Niezgodny nagłówek"), "Nie rozpoznano"),
            (lambda b: b["Klienci"].__setitem__("D2", None), "nagłówkach godzin"),
            (lambda b: b["Klienci"].__setitem__("E2", 1), "zduplikowane godziny"),
            (lambda b: b["Klienci"].__setitem__("A3", "NIEZNANY_OSD"), "nie występuje"),
            (lambda b: b.create_sheet("Klienci_part_1"), "Niejednoznaczny"),
        )
        for change, error in mutations:
            with self.subTest(error=error):
                book = inline_workbook()
                change(book)
                self.save(book)
                with self.assertRaisesRegex(ValueError, f"bialystok_test.xlsx: .*{error}"):
                    load_gams_workbook(self.path)

    def test_partial_split_is_not_treated_as_inline(self):
        book = inline_workbook()
        book["Klienci"].title = "Klienci_part_1"
        domain = book.create_sheet("Klienci")
        domain.append(["PPE"])
        domain.append(["TEST_A"])
        self.save(book)
        with self.assertRaisesRegex(ValueError, "bialystok_test.xlsx: Niekompletny układ"):
            inspect_workbook_clients(self.path, {"TEST_A"})

    def test_quick_schema_error_contains_filename_and_available_sheets(self):
        book = inline_workbook()
        book["Klienci"]["C2"] = "Nieznana kolumna"
        self.save(book)
        with self.assertRaisesRegex(ValueError, "bialystok_test.xlsx: .*Dostępne arkusze: Klienci, Oddzial, Ograniczenia"):
            inspect_workbook_clients(self.path, {"TEST_A"})

    def test_mixed_eight_branch_portfolio_validates_and_solves_1100_pairs(self):
        root = Path(self.temporary.name)
        args = make_portfolio(root, total=1100, eligible_per_branch=280)
        source = root / "model_bialystok.xlsx"
        book = load_workbook(source)
        del book["Klienci"]
        del book["Klienci_part_2"]
        book["Klienci_part_1"].title = "Klienci"
        book["Klienci"]["C2"] = "dataCET"
        book.save(source)
        book.close()
        with patch("representative_clients.portfolio_cli.solve_selection") as solve, redirect_stdout(io.StringIO()):
            self.assertEqual(main([*args, "--validate-full"]), 0)
            solve.assert_not_called()
        output = root / "wynik.xlsx"
        with redirect_stdout(io.StringIO()) as log:
            self.assertEqual(main([*args, "--output", str(output), "--time-limit", "10"]), 0)
        self.assertIn("model_bialystok.xlsx: układ: profile w Klienci", log.getvalue())
        self.assertIn("model_lublin.xlsx: układ: Klienci_part_1 + Klienci_part_2", log.getvalue())
        report = load_workbook(output, read_only=True, data_only=True)
        try:
            pairs = list(report["Zastepcy"].iter_rows(min_row=2, values_only=True))
            self.assertEqual(len(pairs), 1100)
            self.assertEqual(len({row[4] for row in pairs}), 1100)
            self.assertFalse({row[2] for row in pairs} & {row[4] for row in pairs})
        finally:
            report.close()


if __name__ == "__main__":
    unittest.main()
