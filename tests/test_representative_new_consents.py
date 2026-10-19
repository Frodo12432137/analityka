from contextlib import redirect_stderr, redirect_stdout
import io
import hashlib
import json
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch

from openpyxl import Workbook, load_workbook

from representative_clients.allocation_plan import create_plan, load_plan
from representative_clients.consent import BRANCH_ORDER, load_consent_workbook
from representative_clients.launcher_options import build_run_arguments, TIME_MODES
from representative_clients.portfolio_cli import main
from representative_clients.portfolio_data import allocate_branch_targets
from tests.test_representative_launcher import make_portfolio


NEW_HEADERS = ["Numer PPE", "Właściwy PPE (w przypadku błędów w kolumnie A)"]
OLD_HEADERS = ["Kod PP Aktualny", "OSD Nazwa - PP Aktualny"]


def workbook(path, rows, headers=NEW_HEADERS):
    book = Workbook()
    book.active.append(headers)
    for row in rows:
        book.active.append(row)
    book.save(path)
    book.close()


def records(sheet):
    rows = sheet.iter_rows(values_only=True)
    headers = next(rows)
    return [dict(zip(headers, row)) for row in rows]


class NewConsentTests(unittest.TestCase):
    def test_corrected_code_wins_and_old_list_never_expands_current_consents(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            current, old = root / "new.xlsx", root / "old.xlsx"
            workbook(old, [["001", "Białystok"], ["002", "Lublin"], ["003", "Warszawa"]], OLD_HEADERS)
            workbook(current, [["001", "002"], [123, "001"], [None, "001"], [" 002 ", "  "], [None, None]])
            index = load_consent_workbook(current, old)
            self.assertEqual(index.by_branch["bialystok"], frozenset({"001"}))
            self.assertEqual(index.by_branch["lublin"], frozenset({"002"}))
            self.assertFalse(index.by_branch["warszawa"])
            self.assertEqual((index.unique_count, index.source_rows, index.duplicate_rows, index.corrected_rows), (2, 4, 2, 3))

    def test_unmatched_correction_never_falls_back_to_original(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            workbook(root / "old.xlsx", [["001", "Lublin"]], OLD_HEADERS)
            workbook(root / "new.xlsx", [["001", "999"]])
            with self.assertRaisesRegex(ValueError, "Brak przypisania OSD"):
                load_consent_workbook(root / "new.xlsx", root / "old.xlsx")

    def test_invalid_effective_code_and_conflicting_mapping_stop_import(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            old, current = root / "old.xlsx", root / "new.xlsx"
            workbook(old, [["001", "Lublin"]], OLD_HEADERS)
            for row in ([590000000000000001, None], ["001", 590000000000000001], ["001", '=A2'], ['=B2', None]):
                with self.subTest(row=row):
                    workbook(current, [row])
                    with self.assertRaisesRegex(ValueError, "tekstem"):
                        load_consent_workbook(current, old)
            workbook(old, [["001", "Lublin"], ["001", "Warszawa"]], OLD_HEADERS)
            workbook(current, [["001", None]])
            with self.assertRaisesRegex(ValueError, "dwóch OSD"):
                load_consent_workbook(current, old)

    def test_new_layout_requires_mapping_and_nonempty_list(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            workbook(root / "old.xlsx", [["001", "Lublin"]], OLD_HEADERS)
            workbook(root / "new.xlsx", [["001", None]])
            with self.assertRaisesRegex(ValueError, "--osd-mapping"):
                load_consent_workbook(root / "new.xlsx")
            workbook(root / "new.xlsx", [])
            with self.assertRaisesRegex(ValueError, "żadnego PPE"):
                load_consent_workbook(root / "new.xlsx", root / "old.xlsx")


class SmallPortfolioTests(unittest.TestCase):
    def test_existing_version_one_plan_remains_readable(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            domains = {key: {f"{key}_{i}" for i in range(128)} for key in BRANCH_ORDER}
            sources = {key: [(root / f"{key}.xlsx", key)] for key in BRANCH_ORDER}
            plan = create_plan(domains, 500, sources, root / "old.xlsx")
            plan["version"] = 1
            for key in ("with_reserves", "osd_mapping", "all_consented", "plan_id"):
                del plan[key]
            plan["plan_id"] = hashlib.sha256(json.dumps(plan, sort_keys=True).encode()).hexdigest()[:16]
            path = root / "legacy_plan.json"
            path.write_text(json.dumps(plan))
            self.assertEqual(load_plan(path), plan)

    def test_capacity_depends_on_explicit_reserve_mode(self):
        counts = {"bialystok": 270, "lublin": 0}
        self.assertEqual(allocate_branch_targets(counts, 270, with_reserves=False), counts)
        with self.assertRaises(ValueError):
            allocate_branch_targets(counts, 270)
        with self.assertRaises(ValueError):
            allocate_branch_targets(counts, 271, with_reserves=False)

    def test_full_270_point_run_uses_new_list_and_reports_quality_without_reserves(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            args = make_portfolio(root, total=270, eligible_per_branch=34)
            # 272 punktów w starym pliku; tylko 270 w aktualnej liście.
            codes = [f"TEST{branch:02d}{index:05d}" for branch in range(8) for index in range(34)][:-2]
            current, old = root / "new.xlsx", root / "zgody.xlsx"
            workbook(current, [[123, codes[0]], *[[code, None] for code in codes[1:]], [codes[0], None]])
            args[args.index("--consents") + 1] = str(current)
            output, plan = root / "report.xlsx", root / "plan.json"
            args += ["--osd-mapping", str(old), "--no-reserves", "--groups", "1",
                     "--allocation-plan", str(plan), "--output", str(output), "--time-limit", "10", "--mip-gap", "0"]
            with redirect_stdout(io.StringIO()), patch("representative_clients.portfolio_cli.choose_unique_reserves") as reserves:
                self.assertEqual(main(args), 0)
                reserves.assert_not_called()
            self.assertFalse(load_plan(plan)["with_reserves"])
            book = load_workbook(output, data_only=True)
            try:
                chosen = {row["klient"] for row in records(book["Wybrani"])}
                self.assertEqual(chosen, set(codes))
                self.assertEqual(len(records(book["Reprezentatywnosc"])), 8)
                self.assertTrue(all(row["srednie_MAE_model"] is not None for row in records(book["Podsumowanie"])))
                self.assertTrue(all(row["zastepcy"] == 0 for row in records(book["Podsumowanie"])))
                self.assertEqual(len(records(book["Pokrycie_OSD"])), 8)
                mail = " ".join(str(row[0]) for row in book["Tekst_do_maila"].iter_rows(values_only=True))
                self.assertIn("Dobór zastępców wyłączony", mail)
                self.assertNotIn("NIE JEST kompletny", mail)
            finally:
                book.close()
            # Zapisany plan przywraca ścieżkę mapowania, ale chroni tryb zastępców.
            with redirect_stdout(io.StringIO()):
                self.assertEqual(main(["--allocation-plan", str(plan), "--consents", str(current),
                                       "--no-reserves", "--validate-only"]), 0)
            with redirect_stderr(io.StringIO()), self.assertRaises(SystemExit):
                main(["--allocation-plan", str(plan), "--consents", str(current),
                      "--total-clients", "270", "--validate-only"])

    def test_empty_osd_is_explicitly_unassessed_in_report(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            args = make_portfolio(root, total=7, eligible_per_branch=1)
            workbook(root / "new.xlsx", [[f"TEST{branch:02d}00000", None] for branch in range(7)])
            old = args[args.index("--consents") + 1]
            args[args.index("--consents") + 1] = str(root / "new.xlsx")
            output = root / "report.xlsx"
            with redirect_stdout(io.StringIO()):
                self.assertEqual(main([*args, "--osd-mapping", old, "--no-reserves",
                                       "--allocation-plan", str(root / "plan.json"),
                                       "--output", str(output)]), 0)
            self.assertEqual(load_plan(root / "plan.json")["branches"]["zamosc"]["target"], 0)
            book = load_workbook(output, data_only=True)
            try:
                missing = next(row for row in records(book["Pokrycie_OSD"]) if row["OSD"] == "Zamość")
                self.assertIn("nie jest oceniony", missing["status"])
                mail = " ".join(str(row[0]) for row in book["Tekst_do_maila"].iter_rows(values_only=True))
                self.assertIn("Nie oceniono pełnego portfela", mail)
                self.assertEqual(len(records(book["Wybrani"])), 7)
            finally:
                book.close()

    def test_mapping_file_cannot_be_overwritten_by_output(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            args = make_portfolio(root, total=8, eligible_per_branch=2)
            mapping = root / "mapping.xlsx"
            workbook(mapping, [["001", "Lublin"]], OLD_HEADERS)
            original = mapping.read_bytes()
            with redirect_stderr(io.StringIO()), self.assertRaises(SystemExit):
                main([*args, "--osd-mapping", str(mapping), "--output", str(mapping), "--overwrite"])
            self.assertEqual(mapping.read_bytes(), original)

    def test_launcher_uses_excel_count_without_reading_manual_target(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            current, old, source = (root / name for name in ("new.xlsx", "old.xlsx", "model.xlsx"))
            for path in (current, old, source):
                path.touch()
            arguments, _ = build_run_arguments(
                sources=[source], consents=current, destination=root, plan_path=root / "plan.json",
                total="Automatycznie z Excela", branches=["Białystok"], mode="validate", time_mode=TIME_MODES[1],
                osd_mapping=old, with_reserves=True, groups=1,
            )
            self.assertIn("--no-reserves", arguments)
            self.assertEqual(arguments[arguments.index("--osd-mapping") + 1], str(old.resolve()))
            self.assertNotIn("--total-clients", arguments)
            self.assertIn("--all-consented", arguments)
            self.assertEqual(arguments[arguments.index("--groups") + 1], "1")

    def test_new_mapping_automatically_uses_actual_count_and_preserves_error_metrics(self):
        for per_branch in (2, 3):
            with self.subTest(per_branch=per_branch), tempfile.TemporaryDirectory() as directory:
                root = Path(directory)
                args = make_portfolio(root, eligible_per_branch=per_branch)[:-2]
                codes = [f"TEST{branch:02d}{i:05d}" for branch in range(8) for i in range(per_branch)]
                workbook(root / "new.xlsx", [[code, None] for code in codes])
                old = args[args.index("--consents") + 1]
                args[args.index("--consents") + 1] = str(root / "new.xlsx")
                output = root / "automatic.xlsx"
                args += ["--osd-mapping", old, "--groups", "1", "--mip-gap", "0"]
                with redirect_stdout(io.StringIO()), patch("representative_clients.portfolio_cli.choose_unique_reserves") as reserves:
                    self.assertEqual(main([*args, "--output", str(output)]), 0)
                    reserves.assert_not_called()
                    self.assertEqual(main([*args, "--total-clients", str(len(codes)), "--no-reserves",
                                           "--output", str(root / "manual.xlsx")]), 0)
                automatic = load_workbook(output, data_only=True)
                manual = load_workbook(root / "manual.xlsx", data_only=True)
                try:
                    self.assertEqual({row["klient"] for row in records(automatic["Wybrani"])}, set(codes))
                    self.assertEqual(records(automatic["Reprezentatywnosc"]), records(manual["Reprezentatywnosc"]))
                    self.assertEqual(records(automatic["Bilans"]), records(manual["Bilans"]))
                finally:
                    automatic.close()
                    manual.close()

    def test_automatic_count_above_1100_round_trips_plan_and_selected_branch(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            args = make_portfolio(root, eligible_per_branch=140)[:-2]
            plan_path = root / "auto.json"
            with redirect_stdout(io.StringIO()):
                self.assertEqual(main([*args, "--all-consented", "--allocation-plan", str(plan_path), "--plan-only"]), 0)
                plan = load_plan(plan_path)
                self.assertEqual(plan["total_clients"], 1120)
                self.assertTrue(plan["all_consented"])
                self.assertFalse(plan["with_reserves"])
                self.assertEqual(main(["--consents", str(root / "zgody.xlsx"), "--allocation-plan", str(plan_path),
                                       "--branches", "Lublin", "--validate-only"]), 0)

    def test_automatic_mode_does_not_reuse_plan_with_smaller_selection(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            args = make_portfolio(root, total=8, eligible_per_branch=3)
            plan = root / "subset.json"
            with redirect_stdout(io.StringIO()):
                self.assertEqual(main([*args, "--no-reserves", "--allocation-plan", str(plan), "--plan-only"]), 0)
            with redirect_stderr(io.StringIO()) as errors, self.assertRaises(SystemExit):
                main(["--consents", str(root / "zgody.xlsx"), "--allocation-plan", str(plan),
                      "--all-consented", "--validate-only"])
            self.assertIn("ogranicza liczbę PPE", errors.getvalue())


if __name__ == "__main__":
    unittest.main()
