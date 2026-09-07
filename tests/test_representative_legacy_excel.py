from __future__ import annotations

import tempfile
import unittest
from pathlib import Path

import numpy as np
from openpyxl import Workbook

from representative_clients.legacy_excel import load_gams_workbook


class LegacyExcelHardeningTests(unittest.TestCase):
    def _base_workbook(self) -> Workbook:
        workbook = Workbook()
        workbook.remove(workbook.active)

        branch = workbook.create_sheet("Oddzial")
        branch.append([None, None, None, None])
        branch.append(["Oddzial", "Data", 1, 2])

        clients = workbook.create_sheet("Klienci")
        clients.append([None])
        clients.append(["A"])
        clients.append(["B"])

        for name in ("Klienci_part_1", "Klienci_part_2"):
            sheet = workbook.create_sheet(name)
            sheet.append([None, None, None, None, None])
            sheet.append(["Oddzial", "Klient", "Data", 1, 2])

        restrictions = workbook.create_sheet("Ograniczenia")
        restrictions.append([None, None, None, None, None, None])
        return workbook

    def _save_and_load(self, workbook: Workbook, branch: str | None = None):
        temporary_directory = tempfile.TemporaryDirectory()
        self.addCleanup(temporary_directory.cleanup)
        path = Path(temporary_directory.name) / "model.xlsx"
        workbook.save(path)
        return load_gams_workbook(path, branch=branch)

    def test_uses_global_date_hour_cartesian_product_for_selected_branch(self) -> None:
        workbook = self._base_workbook()
        workbook["Oddzial"].append(["OSD_A", "2026-01-01", 10.0, 20.0])
        workbook["Oddzial"].append(["OSD_B", "2026-01-02", 30.0, 40.0])
        workbook["Klienci_part_1"].append(
            ["OSD_A", "A", "2026-01-01", 1.0, 2.0]
        )
        # This date exists only in OSD_B's target, but it belongs to the
        # global s_Data_Hist domain and is therefore valid for OSD_A profiles.
        workbook["Klienci_part_1"].append(
            ["OSD_A", "A", "2026-01-02", 3.0, 4.0]
        )
        workbook["Klienci_part_1"].append(
            ["OSD_B", "B", "2026-01-02", 5.0, 6.0]
        )

        data = self._save_and_load(workbook, branch="OSD_A")

        self.assertEqual(
            data.periods,
            (
                ("2026-01-01", 1),
                ("2026-01-01", 2),
                ("2026-01-02", 1),
                ("2026-01-02", 2),
            ),
        )
        np.testing.assert_allclose(data.branch_balance, [10.0, 20.0, 0.0, 0.0])
        np.testing.assert_allclose(
            data.client_balance,
            [[1.0, 0.0], [2.0, 0.0], [3.0, 0.0], [4.0, 0.0]],
        )

    def test_multibranch_workbook_requires_explicit_branch(self) -> None:
        workbook = self._base_workbook()
        workbook["Oddzial"].append(["OSD_A", "2026-01-01", 1.0, 2.0])
        workbook["Oddzial"].append(["OSD_B", "2026-01-01", 3.0, 4.0])

        with self.assertRaisesRegex(ValueError, "podaj --branch"):
            self._save_and_load(workbook)

    def test_intentional_blank_numeric_cells_still_mean_zero(self) -> None:
        workbook = self._base_workbook()
        workbook["Oddzial"].append(["OSD_A", "2026-01-01", 10.0, None])
        workbook["Klienci_part_1"].append(
            ["OSD_A", "A", "2026-01-01", None, 2.0]
        )

        data = self._save_and_load(workbook)

        np.testing.assert_allclose(data.branch_balance, [10.0, 0.0])
        np.testing.assert_allclose(data.client_balance, [[0.0, 0.0], [2.0, 0.0]])

    def test_accepts_accented_oddzial_sheet_seen_in_real_workbook(self) -> None:
        workbook = self._base_workbook()
        workbook["Oddzial"].title = "Oddział"
        workbook["Oddział"].append(
            ["PGED_Skarzysko_Kamienna", "2025-03-11", 1.0, 2.0]
        )
        workbook["Klienci_part_1"].append(
            ["PGED_Skarzysko_Kamienna", "A", "2025-03-11", 0.1, -0.2]
        )

        data = self._save_and_load(workbook)

        self.assertEqual(data.branch, "PGED_Skarzysko_Kamienna")
        np.testing.assert_allclose(data.branch_balance, [1.0, 2.0])
        np.testing.assert_allclose(data.client_balance[:, 0], [0.1, -0.2])

    def test_formula_without_cached_result_is_not_silently_zeroed(self) -> None:
        workbook = self._base_workbook()
        workbook["Oddzial"].append(["OSD_A", "2026-01-01", 10.0, "=2*10"])

        with self.assertRaisesRegex(ValueError, "formuły bez zapisanej wartości"):
            self._save_and_load(workbook)

    def test_unknown_client_in_another_branch_is_rejected_like_gams_domain(self) -> None:
        workbook = self._base_workbook()
        workbook["Oddzial"].append(["OSD_A", "2026-01-01", 1.0, 2.0])
        workbook["Oddzial"].append(["OSD_B", "2026-01-01", 3.0, 4.0])
        workbook["Klienci_part_1"].append(
            ["OSD_B", "NIEZNANY", "2026-01-01", 1.0, 2.0]
        )

        with self.assertRaisesRegex(ValueError, "nie występuje w arkuszu Klienci"):
            self._save_and_load(workbook, branch="OSD_A")

    def test_duplicate_profile_row_is_rejected(self) -> None:
        workbook = self._base_workbook()
        workbook["Oddzial"].append(["OSD_A", "2026-01-01", 1.0, 2.0])
        workbook["Klienci_part_1"].append(
            ["OSD_A", "A", "2026-01-01", 1.0, 2.0]
        )
        workbook["Klienci_part_1"].append(
            ["OSD_A", "A", "2026-01-01", 3.0, 4.0]
        )

        with self.assertRaisesRegex(ValueError, "duplikat klienta"):
            self._save_and_load(workbook)

    def test_corrupted_xlsx_has_readable_error(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "uszkodzony.xlsx"
            path.write_bytes(b"to nie jest skoroszyt")

            with self.assertRaisesRegex(ValueError, "poprawnego skoroszytu XLSX"):
                load_gams_workbook(path)

    def test_numeric_ppe_is_rejected_before_excel_can_hide_corruption(self) -> None:
        workbook = self._base_workbook()
        workbook["Klienci"]["A2"] = 590000000000000001
        workbook["Oddzial"].append(["OSD_A", "2026-01-01", 1.0, 2.0])

        with self.assertRaisesRegex(ValueError, "PPE musi być zapisany.*jako tekst"):
            self._save_and_load(workbook)


if __name__ == "__main__":
    unittest.main()
