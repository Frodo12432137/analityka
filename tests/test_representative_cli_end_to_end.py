from __future__ import annotations

import io
import tempfile
import unittest
from contextlib import redirect_stdout
from dataclasses import replace
from pathlib import Path
from unittest.mock import patch

from openpyxl import Workbook, load_workbook

from representative_clients.cli import main as cli_main
from representative_clients.legacy_excel import load_gams_workbook
from representative_clients.model import SelectionConfig
from representative_clients.solver import solve_selection


class CliEndToEndTests(unittest.TestCase):
    @staticmethod
    def _write_input(path: Path) -> None:
        workbook = Workbook()
        workbook.remove(workbook.active)

        branch = workbook.create_sheet("Oddzial")
        branch.append([None, None, None, None])
        branch.append(["Oddzial", "Data", 1, 2])
        branch.append(["PGED_TEST", "2026-01-01", 2.0, 4.0])

        clients = workbook.create_sheet("Klienci")
        clients.append([None])
        clients.append(["A"])

        for name, values in (
            ("Klienci_part_1", (1.0, 2.0)),
            ("Klienci_part_2", (0.0, 0.0)),
        ):
            sheet = workbook.create_sheet(name)
            sheet.append([None, None, None, None, None])
            sheet.append(["Oddzial", "Klient", "Data", 1, 2])
            sheet.append(["PGED_TEST", "A", "2026-01-01", *values])

        restrictions = workbook.create_sheet("Ograniczenia")
        restrictions.append([None] * 6)
        restrictions.append([None] * 6)
        workbook.save(path)

    def test_real_parser_solver_and_export_pipeline(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            source = root / "model.xlsx"
            output = root / "wynik.xlsx"
            self._write_input(source)

            with redirect_stdout(io.StringIO()):
                return_code = cli_main(
                    [
                        "--input",
                        str(source),
                        "--branch",
                        "PGED_TEST",
                        "--output",
                        str(output),
                        "--groups",
                        "1",
                        "--max-clients",
                        "1",
                        "--max-share",
                        "10",
                        "--mip-gap",
                        "0",
                        "--time-limit",
                        "30",
                    ]
                )

            self.assertEqual(return_code, 0)
            self.assertTrue(output.exists())
            workbook = load_workbook(output, read_only=True, data_only=True)
            try:
                self.assertEqual(
                    workbook.sheetnames,
                    [
                        "Wyniki",
                        "Wybrani",
                        "Bilans",
                        "Istotnosc_dzienna",
                        "Istotnosc_all",
                        "Diagnostyka",
                        "Konfiguracja",
                    ],
                )
                self.assertEqual(workbook["Wybrani"]["C2"].value, "A")
                self.assertAlmostEqual(workbook["Wybrani"]["D2"].value, 2.0)
                self.assertAlmostEqual(workbook["Bilans"]["F2"].value, 2.0)
            finally:
                workbook.close()

    def test_feasible_but_nonoptimal_result_is_exported_with_exit_code_two(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            source = root / "model.xlsx"
            output = root / "wynik_limit_czasu.xlsx"
            self._write_input(source)
            data = load_gams_workbook(source, branch="PGED_TEST")
            config = SelectionConfig(
                max_clients_per_group=1,
                group_count=1,
                max_contribution_share=10.0,
                time_limit_seconds=30.0,
                mip_relative_gap=0.0,
            )
            solved = solve_selection(data, config)
            limited_group = replace(
                solved.groups[0],
                status=1,
                status_text="Time limit reached",
                optimal=False,
                mip_gap=0.25,
            )
            limited = replace(solved, groups=(limited_group,))

            with patch(
                "representative_clients.cli.solve_selection", return_value=limited
            ), redirect_stdout(io.StringIO()):
                return_code = cli_main(
                    [
                        "--input",
                        str(source),
                        "--branch",
                        "PGED_TEST",
                        "--output",
                        str(output),
                        "--groups",
                        "1",
                        "--max-clients",
                        "1",
                        "--max-share",
                        "10",
                        "--mip-gap",
                        "0",
                        "--time-limit",
                        "30",
                    ]
                )

            self.assertEqual(return_code, 2)
            workbook = load_workbook(output, read_only=True, data_only=True)
            try:
                self.assertEqual(workbook["Diagnostyka"]["H2"].value, False)
                self.assertAlmostEqual(workbook["Diagnostyka"]["J2"].value, 0.25)
            finally:
                workbook.close()


if __name__ == "__main__":
    unittest.main()
