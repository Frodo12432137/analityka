"""Regresje na danych syntetycznych, bez identyfikatorów i wyników użytkownika."""
from contextlib import redirect_stdout
from dataclasses import replace
import io
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch
from types import SimpleNamespace

import numpy as np
from openpyxl import Workbook, load_workbook

from representative_clients.model import SelectionConfig, SelectionData
from representative_clients.portfolio_cli import main
from representative_clients.portfolio_data import inspect_workbook_clients
from representative_clients.portfolio_report import _style_workbook
from representative_clients.replacement import choose_unique_reserves
from representative_clients.solver import (
    NoFeasibleSolutionError, selection_weight_bounds, solve_group, solve_selection,
)
from tests.test_representative_new_consents import records


class WeightPopulationTests(unittest.TestCase):
    @staticmethod
    def case():
        data = SelectionData(
            branch="PGED_Bialystok",
            client_ids=tuple(f"TEST{i}" for i in range(20)),
            periods=(("2026-01-01", 1), ("2026-01-01", 2)),
            client_balance=np.tile([[1.0], [-0.5]], (1, 20)),
            branch_balance=np.array([10000.0, -5000.0]),
        )
        config = SelectionConfig(
            group_count=2, max_clients_per_group=10,
            exact_clients_per_group=(10, 10), mip_relative_gap=0,
        )
        return data, config

    def test_source_population_restores_scale_and_both_solver_paths_agree(self):
        data, config = self.case()
        old = solve_selection(data, config)
        self.assertAlmostEqual(old.groups[0].objective_value / 15000, 0.99)
        self.assertEqual(old.groups[0].weights_at_upper_bound, 10)
        fixed = replace(config, weight_population_size=4000)
        result = solve_selection(data, fixed)
        self.assertEqual(result.selected_count, 20)
        for group in result.groups:
            np.testing.assert_allclose(group.model_balance, data.branch_balance, atol=1e-7)
            self.assertLess(group.objective_value, 1e-7)
            self.assertEqual(group.weight_upper_bound, 2000)
            self.assertEqual(group.weight_lower_bound, 0.2)
            direct = solve_group(data, fixed, group.group, group.candidate_ids, selection_count=10)
            np.testing.assert_allclose(direct.model_balance, group.model_balance, atol=1e-7)

    def test_population_does_not_remove_contribution_cap(self):
        data, config = self.case()
        result = solve_selection(data, replace(
            config, weight_population_size=4000, max_contribution_share=0.01,
        ))
        for group in result.groups:
            self.assertTrue(all(w <= 75 + 1e-7 for w in group.selected_weights.values()))
            self.assertEqual(group.weights_at_upper_bound, 10)
            self.assertGreater(group.objective_value, 0)

    def test_population_increases_upper_without_excluding_peaked_client(self):
        data = SelectionData(
            branch="SYNTHETIC", client_ids=("TEST_A", "TEST_B"),
            periods=(("2026-01-01", 1), ("2026-01-01", 2)),
            client_balance=np.array([[100.0, 1.0], [1.0, 1.0]]),
            branch_balance=np.array([1000.0, 1000.0]),
        )
        legacy = SelectionConfig(
            group_count=1, max_clients_per_group=2,
            exact_clients_per_group=(2,), mip_relative_gap=0,
        )
        # Dawny sposób liczenia minimum z populacji dawał 50, choć
        # pierwszy klient ma indywidualny limit wagi 2: model był niewykonalny.
        config = replace(legacy, weight_population_size=1000)
        self.assertEqual(selection_weight_bounds(data, config), (0.1, 2500))
        old = solve_selection(data, legacy)
        fixed = solve_selection(data, config)
        self.assertEqual(fixed.selected_count, 2)
        group = fixed.groups[0]
        self.assertLessEqual(group.selected_weights["TEST_A"], 2 + 1e-7)
        self.assertLessEqual(group.objective_value, old.groups[0].objective_value + 1e-7)
        # Jawna mała waga wymagana, zgodna z poprzednim minimum, jest zachowana.
        required = replace(data, required_weights={"TEST_A": 1.0})
        self.assertEqual(solve_selection(required, config).groups[0].selected_weights["TEST_A"], 1.0)
        direct = solve_group(data, config, "w1", data.client_ids, selection_count=2)
        self.assertAlmostEqual(direct.objective_value, group.objective_value)

    def test_exact_group_selection_relaxes_weight_floor_not_hard_caps(self):
        data = SelectionData(
            branch="SYNTHETIC", client_ids=("TEST_A", "TEST_B"),
            periods=(("2026-01-01", 1),),
            client_balance=np.array([[1.0, 1000.0]]),
            branch_balance=np.array([100.0]),
        )
        config = SelectionConfig(
            group_count=2, max_clients_per_group=1, exact_clients_per_group=(1, 1),
            weight_population_size=1000, mip_relative_gap=0,
        )
        result = solve_selection(data, config)
        self.assertEqual(result.selected_count, 2)
        self.assertEqual(result.groups[0].weight_lower_bound, 0.2)
        self.assertAlmostEqual(result.groups[1].weight_lower_bound, 0.02, places=9)
        self.assertEqual(result.groups[1].weight_lower_nominal, 0.2)
        self.assertIn("minimalna waga obniżona", result.groups[1].status_text)
        self.assertLessEqual(result.groups[1].selected_weights["TEST_B"], 0.02 + 1e-7)
        direct = solve_group(data, config, "w2", ("TEST_B",), selection_count=1)
        self.assertAlmostEqual(direct.weight_lower_bound, result.groups[1].weight_lower_bound)

    def test_all_candidates_selected_even_with_zero_contribution_caps(self):
        data = SelectionData(
            branch="SYNTHETIC",
            client_ids=("ZERO_A", "ZERO_B", "SMALL_C", "SMALL_D"),
            periods=(("2026-01-01", 1), ("2026-01-01", 2)),
            client_balance=np.array([[1, 2, 5000, -5000],
                                     [1, 2, 5000, -5000]], dtype=float),
            branch_balance=np.array([0.0, 0.0]),
        )
        config = SelectionConfig(
            group_count=2, max_clients_per_group=2, exact_clients_per_group=(2, 2),
            weight_population_size=100_000, mip_relative_gap=0,
        )
        result = solve_selection(data, config)
        self.assertEqual(result.selected_count, 4)
        self.assertTrue(all(g.weight_lower_bound == 0 for g in result.groups))
        self.assertTrue(all(g.weight_lower_nominal == 0.2 for g in result.groups))
        self.assertTrue(all(all(w == 0 for w in g.selected_weights.values()) for g in result.groups))
        np.testing.assert_allclose(result.groups[0].model_balance, [0, 0])

    def test_adaptive_floor_retains_normal_non_exact_mode(self):
        data = SelectionData(
            branch="SYNTHETIC",
            client_ids=("A", "B"),
            periods=(("2026-01-01", 1),),
            client_balance=np.array([[1.0, 1000.0]]),
            branch_balance=np.array([100.0]),
        )
        config = SelectionConfig(
            group_count=1, max_clients_per_group=2,
            weight_population_size=1000, mip_relative_gap=0,
        )
        result = solve_selection(data, config)
        self.assertEqual(result.groups[0].weight_lower_bound, 0.1)
        self.assertEqual(result.groups[0].weight_lower_nominal, 0.1)

    def test_required_weight_above_contribution_cap_is_still_rejected(self):
        data = SelectionData(
            branch="SYNTHETIC",
            client_ids=("A",),
            periods=(("2026-01-01", 1),),
            client_balance=np.array([[1000.0]]),
            branch_balance=np.array([100.0]),
            required_weights={"A": 1.0},
        )
        config = SelectionConfig(
            group_count=1, max_clients_per_group=1, exact_clients_per_group=(1,),
            weight_population_size=1000,
        )
        with self.assertRaisesRegex(NoFeasibleSolutionError, "Wymagana waga"):
            solve_selection(data, config)

    def test_weight_floor_adapts_across_distinct_group_shapes(self):
        data = SelectionData(
            branch="PGED_Lublin",
            client_ids=("A", "B", "C", "D", "E", "F"),
            periods=(("2026-01-01", 1), ("2026-01-01", 2)),
            client_balance=np.array([[1, 2, 100, 200, 1000, 2000],
                                     [1, 2, 100, 200, 1000, 2000]], dtype=float),
            branch_balance=np.array([100.0, 100.0]),
        )
        config = SelectionConfig(
            group_count=3, max_clients_per_group=2, exact_clients_per_group=(2, 2, 2),
            weight_population_size=50000, mip_relative_gap=0,
        )
        result = solve_selection(data, config)
        self.assertEqual(result.selected_count, 6)
        for group in result.groups:
            self.assertEqual(len(group.selected_weights), 2)
            self.assertGreaterEqual(group.weight_lower_bound, 0)
            self.assertLessEqual(group.weight_lower_bound, group.weight_lower_nominal)
            self.assertLessEqual(group.max_constraint_violation, 1e-4)

    def test_no_solution_message_distinguishes_timeout_and_infeasibility(self):
        data, config = self.case()
        for status, expected in ((1, "limit czasu lub iteracji"), (2, "zadanie niewykonalne")):
            with self.subTest(status=status):
                fake = SimpleNamespace(status=status, x=None, message="synthetic solver message")
                with patch("scipy.optimize.milp", return_value=fake):
                    with self.assertRaisesRegex(NoFeasibleSolutionError, expected):
                        solve_selection(data, config)

    def test_reserves_can_match_primary_weights_above_legacy_limit(self):
        data, config = self.case()
        data = replace(data, branch_balance=data.branch_balance * 10)
        config = replace(config, group_count=1, exact_clients_per_group=(10,),
                         weight_population_size=40000)
        result = solve_selection(data, config)
        self.assertLess(result.groups[0].objective_value, 1e-7)
        self.assertTrue(any(w > 3000 for w in result.groups[0].selected_weights.values()))
        replacements = choose_unique_reserves(data, result)
        self.assertEqual(len(replacements.matches), 10)
        self.assertTrue(all(match.l1_error < 1e-7 for match in replacements.matches))
        self.assertLess(replacements.diagnostics[0].l1_after_replacement, 1e-7)
        limited = choose_unique_reserves(data, result, maximum_weight=3000)
        self.assertGreater(limited.diagnostics[0].l1_after_replacement, 1)

    def test_invalid_population_is_rejected_and_default_is_legacy(self):
        data, config = self.case()
        for value in (0, -1, True, 1.5, float("inf")):
            with self.subTest(value=value), self.assertRaises(ValueError):
                replace(config, weight_population_size=value).validate()
        with self.assertRaisesRegex(ValueError, "mniejsza"):
            solve_selection(data, replace(config, weight_population_size=19))
        result = solve_selection(data, config)
        self.assertEqual(result.groups[0].weight_upper_bound, 10)

    @staticmethod
    def source_book():
        book = Workbook()
        clients = book.active
        clients.title = "Klienci"
        clients.append(["PPE"])
        for i in range(4000):
            clients.append([f"TEST{i}"])
        target = book.create_sheet("Oddzial")
        target.append(["Bilans"])
        target.append(["Oddzial", "Data", 1, 2])
        target.append(["PGED_Bialystok", "2026-01-01", 10000, -5000])
        for name in ("Klienci_part_1", "Klienci_part_2"):
            sheet = book.create_sheet(name)
            sheet.append(["Profile"])
            sheet.append(["Oddzial", "Klient", "Data", 1, 2])
        for i in range(20):
            book["Klienci_part_1"].append(["PGED_Bialystok", f"TEST{i}", "2026-01-01", 1, -0.5])
        book.create_sheet("Ograniczenia").append(["Wykluczeni", None, None, None, "Wymagani", "Waga"])
        return book

    def test_portfolio_uses_union_before_consents_including_skipped_part(self):
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            source = root / "source.xlsx"
            book = self.source_book()
            book.save(source)
            # Dodatkowa część: duplikaty populacji + nowe PPE, brak kandydatów.
            # Nie może ani zmniejszyć populacji, ani podwoić wspólnych PPE.
            extra = root / "extra.xlsx"
            book["Klienci"].delete_rows(2, 20)
            for i in range(4000, 4100):
                book["Klienci"].append([f"TEST{i}"])
            book["Klienci_part_1"].delete_rows(3, 20)
            book.save(extra)
            book.close()
            consents = root / "consents.xlsx"
            book = Workbook()
            book.active.append(["Kod PP Aktualny", "OSD Nazwa - PP Aktualny"])
            for i in range(20):
                book.active.append([f"TEST{i}", "PGED_Bialystok"])
            book.save(consents)
            book.close()
            output = root / "result.xlsx"
            with redirect_stdout(io.StringIO()):
                code = main([
                    "--input", str(source), "--input", str(extra),
                    "--consents", str(consents), "--all-consented",
                    "--allow-partial", "--groups", "2", "--mip-gap", "0",
                    "--output", str(output),
                ])
            self.assertEqual(code, 0)
            book = load_workbook(output, data_only=True)
            try:
                selected = records(book["Wybrani"])
                self.assertEqual({row["klient"] for row in selected}, {f"TEST{i}" for i in range(20)})
                self.assertTrue(all(row["populacja_do_wag"] == 4100 for row in records(book["Solver"])))
                self.assertTrue(all(row["waga_max_populacja"] == 2050 for row in records(book["Solver"])))
                self.assertTrue(all(abs(row["blad_wzgledny"]) < 1e-7 for row in records(book["Reprezentatywnosc"])))
                for row in records(book["Bilans"]):
                    self.assertAlmostEqual(row["bilans_model"], row["bilans_historyczny"])
                headers = [cell.value for cell in book["Bilans"][1]]
                absolute = book["Bilans"].cell(2, headers.index("blad_bezwzgledny") + 1)
                self.assertNotIn("%", absolute.number_format)
            finally:
                book.close()

    def test_multibranch_population_is_scoped_for_split_and_inline_layouts(self):
        for inline in (False, True):
            with self.subTest(inline=inline), tempfile.TemporaryDirectory() as folder:
                book = self.source_book()
                book["Oddzial"].append(["PGED_Lublin", "2026-01-01", 20, -10])
                book["Klienci_part_1"].append(["PGED_Lublin", "TEST20", "2026-01-01", 2, -1])
                # Ten sam PPE i inna data nie zwiększają liczebności populacji.
                book["Klienci_part_2"].append(["PGED_Bialystok", "TEST0", "2026-01-02", 1, -0.5])
                if inline:
                    del book["Klienci"]
                    book["Klienci_part_1"].append(["PGED_Bialystok", "TEST0", "2026-01-02", 1, -0.5])
                    del book["Klienci_part_2"]
                    book["Klienci_part_1"].title = "Klienci"
                path = Path(folder) / "source.xlsx"
                book.save(path)
                book.close()
                population, allowed, _, _ = inspect_workbook_clients(
                    path, {"TEST0", "TEST20"}, branch="PGED_Bialystok",
                )
                self.assertEqual(population, {f"TEST{i}" for i in range(20)})
                self.assertEqual(allowed, {"TEST0"})
                population, allowed, _, _ = inspect_workbook_clients(
                    path, {"TEST0", "TEST20"}, branch="PGED_Lublin",
                )
                self.assertEqual(population, {"TEST20"})
                self.assertEqual(allowed, {"TEST20"})


class PercentageFormatTests(unittest.TestCase):
    def test_absolute_errors_are_numbers_and_relative_errors_are_percentages(self):
        book = Workbook()
        sheet = book.active
        sheet.append(["blad_bezwzgledny", "blad_model_minus_historia", "blad_wzgledny",
                      "bias_wzgledny", "mip_gap", "L1", "wymagany"])
        sheet.append([123.45, -123.45, 0.99, -0.99, 0, 123.45, False])
        sheet.append([123, -123, 1, -1, 1, 123, True])
        _style_workbook(book)
        for row in (2, 3):
            for col in (1, 2, 6, 7):
                self.assertNotIn("%", sheet.cell(row, col).number_format)
            for col in (3, 4, 5):
                self.assertEqual(sheet.cell(row, col).number_format, "0.0%")
        self.assertEqual(sheet["A2"].value, 123.45)
        book.close()


if __name__ == "__main__":
    unittest.main()
