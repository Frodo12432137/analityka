from __future__ import annotations

import io
import tempfile
import unittest
from contextlib import ExitStack, redirect_stdout
from pathlib import Path
from unittest.mock import patch

import numpy as np
from openpyxl import Workbook, load_workbook

from representative_clients.grouping import round_robin_groups
from representative_clients.cli import main as cli_main
from representative_clients.legacy_excel import load_gams_workbook
from representative_clients.model import SelectionConfig, SelectionData
from representative_clients.report import write_result_workbook
from representative_clients.solver import (
    contribution_weight_caps,
    gams_weight_bounds,
    solve_selection,
)


class GroupingTests(unittest.TestCase):
    def test_round_robin_preserves_eligible_order(self) -> None:
        groups = round_robin_groups(
            ("A", "B", "C", "D", "E", "F"), {"B"}, group_count=4
        )
        self.assertEqual(groups["w1"], ("A", "F"))
        self.assertEqual(groups["w2"], ("C",))
        self.assertEqual(groups["w3"], ("D",))
        self.assertEqual(groups["w4"], ("E",))

    def test_gams_weight_bounds_use_total_client_count(self) -> None:
        self.assertEqual(gams_weight_bounds(1_000, 50), (2.0, 100.0))

    def test_contribution_cap_is_equivalent_to_all_period_constraints(self) -> None:
        profiles = np.asarray([[2.0, -4.0], [-5.0, 1.0], [1.0, 0.0]])
        target = np.asarray([10.0, -20.0, 30.0])
        caps = contribution_weight_caps(profiles, target, 0.20)
        expected_limit = 0.20 * np.mean(np.abs(target))
        np.testing.assert_allclose(caps, [expected_limit / 5.0, expected_limit / 4.0])


class LegacyWorkbookTests(unittest.TestCase):
    def _make_workbook(self, path: Path) -> None:
        workbook = Workbook()
        workbook.remove(workbook.active)

        branch = workbook.create_sheet("Oddzial")
        branch.append([None, None, None, None])
        branch.append(["Oddzial", "Data", 1, 2])
        branch.append(["PGED_TEST", "2026-01-01", 3.0, 6.0])

        clients = workbook.create_sheet("Klienci")
        clients.append([None])
        clients.append(["A"])
        clients.append(["B"])

        part_1 = workbook.create_sheet("Klienci_part_1")
        part_1.append([None, None, None, None, None])
        part_1.append(["Oddzial", "Klient", "Data", 1, 2])
        part_1.append(["PGED_TEST", "A", "2026-01-01", 1.0, 2.0])
        part_1.append(["PGED_TEST", "B", "2026-01-01", 0.0, 1.0])

        part_2 = workbook.create_sheet("Klienci_part_2")
        part_2.append([None, None, None, None, None])
        part_2.append(["Oddzial", "Klient", "Data", 1, 2])
        part_2.append(["PGED_TEST", "A", "2026-01-01", 0.5, 1.0])
        part_2.append(["PGED_TEST", "B", "2026-01-01", 0.0, 0.0])

        restrictions = workbook.create_sheet("Ograniczenia")
        restrictions.append([None, None, None, None, None, None])
        restrictions.append(["B", None, None, None, "A", 2.0])
        workbook.save(path)

    def test_loads_gams_layout_and_sums_both_profile_parts(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "model.xlsx"
            self._make_workbook(path)
            data = load_gams_workbook(path)

        self.assertEqual(data.branch, "PGED_TEST")
        self.assertEqual(data.client_ids, ("A", "B"))
        self.assertEqual(data.periods, (("2026-01-01", 1), ("2026-01-01", 2)))
        np.testing.assert_allclose(data.branch_balance, [3.0, 6.0])
        np.testing.assert_allclose(data.client_balance, [[1.5, 0.0], [3.0, 1.0]])
        self.assertEqual(data.excluded, frozenset({"B"}))
        self.assertEqual(data.required_weights, {"A": 2.0})


class CliSmokeTests(unittest.TestCase):
    @staticmethod
    def _data() -> SelectionData:
        return SelectionData(
            branch="PGED_TEST",
            client_ids=("A",),
            periods=(("2026-01-01", 1),),
            client_balance=np.asarray([[1.0]]),
            branch_balance=np.asarray([1.0]),
        )

    def test_validate_only_reads_data_but_never_solves_or_writes(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            input_path = Path(directory) / "model.xlsx"
            with ExitStack() as stack:
                load_mock = stack.enter_context(
                    patch(
                        "representative_clients.cli.load_gams_workbook",
                        return_value=self._data(),
                    )
                )
                solve_mock = stack.enter_context(
                    patch("representative_clients.cli.solve_selection")
                )
                write_mock = stack.enter_context(
                    patch("representative_clients.cli.write_result_workbook")
                )
                stack.enter_context(redirect_stdout(io.StringIO()))
                return_code = cli_main(
                    ["--input", str(input_path), "--validate-only"]
                )

            self.assertEqual(return_code, 0)
            load_mock.assert_called_once()
            self.assertEqual(load_mock.call_args.args, (input_path,))
            self.assertIsNone(load_mock.call_args.kwargs["branch"])
            self.assertTrue(callable(load_mock.call_args.kwargs["on_stage"]))
            solve_mock.assert_not_called()
            write_mock.assert_not_called()
            self.assertEqual(list(Path(directory).iterdir()), [])

    def test_existing_output_is_rejected_before_loading_or_solving(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            input_path = root / "model.xlsx"
            output_path = root / "wynik.xlsx"
            output_path.touch()
            with ExitStack() as stack:
                load_mock = stack.enter_context(
                    patch("representative_clients.cli.load_gams_workbook")
                )
                solve_mock = stack.enter_context(
                    patch("representative_clients.cli.solve_selection")
                )
                write_mock = stack.enter_context(
                    patch("representative_clients.cli.write_result_workbook")
                )
                with self.assertRaisesRegex(SystemExit, "już istnieje"):
                    cli_main(
                        [
                            "--input",
                            str(input_path),
                            "--output",
                            str(output_path),
                        ]
                    )

            load_mock.assert_not_called()
            solve_mock.assert_not_called()
            write_mock.assert_not_called()

    def test_same_input_and_output_path_is_rejected(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "model.xlsx"
            path.touch()
            with ExitStack() as stack:
                load_mock = stack.enter_context(
                    patch("representative_clients.cli.load_gams_workbook")
                )
                solve_mock = stack.enter_context(
                    patch("representative_clients.cli.solve_selection")
                )
                write_mock = stack.enter_context(
                    patch("representative_clients.cli.write_result_workbook")
                )
                with self.assertRaisesRegex(SystemExit, "tym samym plikiem"):
                    cli_main(["--input", str(path), "--output", str(path)])

            load_mock.assert_not_called()
            solve_mock.assert_not_called()
            write_mock.assert_not_called()


class SolverTests(unittest.TestCase):
    def test_selects_one_profile_for_exact_fit(self) -> None:
        data = SelectionData(
            branch="PGED_TEST",
            client_ids=("A", "B"),
            periods=(("2026-01-01", 1), ("2026-01-01", 2)),
            client_balance=np.asarray([[1.0, 1.0], [2.0, 0.0]]),
            branch_balance=np.asarray([3.0, 6.0]),
        )
        config = SelectionConfig(
            max_clients_per_group=1,
            group_count=1,
            max_contribution_share=10.0,
            time_limit_seconds=30.0,
            mip_relative_gap=0.0,
        )
        result = solve_selection(data, config)
        group = result.groups[0]
        self.assertEqual(tuple(group.selected_weights), ("A",))
        self.assertAlmostEqual(group.selected_weights["A"], 3.0, places=7)
        self.assertAlmostEqual(group.objective_value, 0.0, places=7)
        np.testing.assert_allclose(group.model_balance, data.branch_balance, atol=1e-7)

    def test_required_client_keeps_fixed_weight(self) -> None:
        data = SelectionData(
            branch="PGED_TEST",
            client_ids=("A", "B"),
            periods=(("2026-01-01", 1), ("2026-01-01", 2)),
            client_balance=np.asarray([[1.0, 0.0], [2.0, 1.0]]),
            branch_balance=np.asarray([2.0, 4.0]),
            required_weights={"A": 2.0},
        )
        config = SelectionConfig(
            max_clients_per_group=1,
            group_count=1,
            max_contribution_share=10.0,
            time_limit_seconds=30.0,
            mip_relative_gap=0.0,
        )
        group = solve_selection(data, config).groups[0]
        self.assertEqual(group.selected_weights, {"A": 2.0})

    def test_writes_auditable_workbook(self) -> None:
        data = SelectionData(
            branch="PGED_TEST",
            client_ids=("A",),
            periods=(("2026-01-01", 1), ("2026-01-01", 2)),
            client_balance=np.asarray([[1.0], [-2.0]]),
            branch_balance=np.asarray([1.0, -2.0]),
        )
        config = SelectionConfig(
            max_clients_per_group=1,
            group_count=1,
            max_contribution_share=10.0,
            time_limit_seconds=30.0,
            mip_relative_gap=0.0,
        )
        result = solve_selection(data, config)
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "wynik.xlsx"
            write_result_workbook(data, result, path)
            workbook = load_workbook(path, read_only=True)
            try:
                self.assertEqual(
                    workbook.sheetnames,
                    [
                        "Wyniki",
                        "Wybrani",
                        "Bilans",
                        "Istotnosc_dzienna",
                        "Istotnosc_all",
                        "Diagnostyka",
                        "Konfiguracja",
                    ],
                )
                self.assertEqual(workbook["Wyniki"]["B3"].value, "w1")
                self.assertEqual(workbook["Wyniki"]["C3"].value, "A")
                self.assertAlmostEqual(workbook["Wyniki"]["D3"].value, 1.0)
                self.assertEqual(workbook["Wyniki"]["J2"].value, "b_historia")
                self.assertEqual(
                    workbook["Wyniki"]["K2"].value, "b_model_w_klienci"
                )
                self.assertEqual(workbook["Wyniki"]["Q2"].value, "Pobor")
                self.assertEqual(workbook["Wyniki"]["X2"].value, "Oddanie")
                self.assertEqual(
                    tuple(cell.value for cell in workbook["Wybrani"][1]),
                    (
                        "oddzial",
                        "grupa",
                        "klient",
                        "waga",
                        "wymagany",
                        "status_solvera",
                        "cel_L1",
                    ),
                )
                self.assertEqual(
                    tuple(cell.value for cell in workbook["Bilans"][1]),
                    (
                        "oddzial",
                        "grupa",
                        "data",
                        "godzina",
                        "bilans_historyczny",
                        "bilans_model",
                        "blad_model_minus_historia",
                        "blad_bezwzgledny",
                    ),
                )
                self.assertEqual(workbook["Wybrani"]["A2"].value, "PGED_TEST")
                self.assertEqual(workbook["Wybrani"]["C2"].value, "A")
                self.assertEqual(workbook["Diagnostyka"]["H2"].value, True)
            finally:
                workbook.close()


if __name__ == "__main__":
    unittest.main()
