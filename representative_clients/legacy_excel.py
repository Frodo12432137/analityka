from __future__ import annotations

import re
import unicodedata
from collections import OrderedDict
from datetime import date, datetime
from pathlib import Path
from posixpath import join as posix_join, normpath as posix_normpath
from typing import AbstractSet, Callable, Dict, List, Optional, Tuple
from xml.etree import ElementTree
from zipfile import BadZipFile, ZipFile

import numpy as np

from .model import Period, SelectionData


_SHEET_ALIASES = {
    "Oddzial": ("Oddzial", "Oddział"),
    "Klienci": ("Klienci",),
    "Klienci_part_1": ("Klienci_part_1",),
    "Klienci_part_2": ("Klienci_part_2",),
    "Ograniczenia": ("Ograniczenia",),
}


def _resolve_sheet_names(workbook: object) -> Dict[str, str]:
    """Rozpoznaje układ dzielony v8 lub profile bezpośrednio w Klienci."""

    available = set(workbook.sheetnames)
    resolved: Dict[str, str] = {}
    missing = []
    for logical_name, aliases in _SHEET_ALIASES.items():
        if logical_name.startswith("Klienci_part_"):
            continue
        matches = [name for name in aliases if name in available]
        if not matches:
            missing.append(" / ".join(aliases))
            continue
        if len(matches) > 1:
            raise ValueError(
                f"Skoroszyt zawiera jednocześnie zakładki {matches!r}. "
                "Pozostaw tylko jeden wariant nazwy."
            )
        resolved[logical_name] = matches[0]
    if missing:
        raise ValueError(
            "Brak wymaganych arkuszy: " + ", ".join(missing)
            + ". Dostępne: " + ", ".join(workbook.sheetnames)
        )
    inline = _has_inline_profile_header(workbook[resolved["Klienci"]])
    parts = [name for name in ("Klienci_part_1", "Klienci_part_2") if name in available]
    if inline and parts:
        raise ValueError(
            "Niejednoznaczny układ: profile występują w Klienci oraz istnieją "
            "arkusze Klienci_part_*. Nie można bezpiecznie ustalić, które źródło czytać."
        )
    if len(parts) == 2:
        resolved.update({name: name for name in parts})
    elif parts:
        raise ValueError(
            "Niekompletny układ dzielony: potrzebne są oba arkusze "
            "Klienci_part_1 i Klienci_part_2. Dostępne: " + ", ".join(workbook.sheetnames)
        )
    elif inline:
        _hours_from_header(workbook[resolved["Klienci"]], 2, 4)
    else:
        raise ValueError(
            "Nie rozpoznano układu źródła. Oczekiwano Klienci_part_1 i Klienci_part_2 "
            "albo profili w Klienci: A2=Oddzial, B2=Klient, C2=dataCET/Data, D2 i dalej=godziny. "
            "Dostępne arkusze: " + ", ".join(workbook.sheetnames)
        )
    return resolved


def _has_inline_profile_header(sheet: object) -> bool:
    header = next(sheet.iter_rows(min_row=2, max_row=2, max_col=3, values_only=True), ())
    def normalize(value):
        text = unicodedata.normalize("NFKD", str(value or "").strip().casefold().replace("ł", "l"))
        return "".join(char for char in text if not unicodedata.combining(char) and char not in " _")
    return len(header) == 3 and (
        normalize(header[0]) in {"oddzial", "osd"}
        and normalize(header[1]) == "klient"
        and normalize(header[2]) in {"datacet", "data", "date"}
    )


def _profile_sheet_names(sheets: Dict[str, str]) -> Tuple[str, ...]:
    if "Klienci_part_1" in sheets:
        return sheets["Klienci_part_1"], sheets["Klienci_part_2"]
    return (sheets["Klienci"],)


def _identifier(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return str(value)
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _client_identifier(value: object, context: str) -> str:
    """PPE musi pozostać tekstem, aby Excel nie uszkodził cyfr ani zer."""

    if value is None or (isinstance(value, str) and not value.strip()):
        return ""
    if not isinstance(value, str):
        raise ValueError(
            f"{context}: identyfikator PPE musi być zapisany w Excelu jako tekst, "
            "nie jako liczba. Liczby mogą utracić zera wiodące lub cyfry po "
            "15. pozycji."
        )
    return value.strip()


def _axis_key(value: object) -> object:
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return int(value)
    if isinstance(value, str):
        return value.strip()
    return value


def _number(value: object, context: str) -> float:
    if value is None or (isinstance(value, str) and not value.strip()):
        return 0.0
    if isinstance(value, str):
        value = value.replace("\u00a0", "").replace(" ", "").replace(",", ".")
    try:
        number = float(value)
    except (TypeError, ValueError) as exc:
        raise ValueError(f"Nieprawidłowa liczba w {context}: {value!r}.") from exc
    if not np.isfinite(number):
        raise ValueError(f"Nieskończona wartość albo NaN w {context}.")
    return number


def _row_values(sheet: object, row_number: int) -> Tuple[object, ...]:
    rows = sheet.iter_rows(
        min_row=row_number, max_row=row_number, values_only=True
    )
    return tuple(next(rows, ()))


def _hours_from_header(
    sheet: object, header_row: int, first_hour_column: int
) -> Tuple[object, ...]:
    row = _row_values(sheet, header_row)
    values = list(row[first_hour_column - 1 :])
    while values and values[-1] is None:
        values.pop()
    hours = tuple(_axis_key(value) for value in values)
    if not hours or any(value in (None, "") for value in hours):
        raise ValueError(
            f"Arkusz {sheet.title!r}: brak lub luka w nagłówkach godzin."
        )
    if len(set(hours)) != len(hours):
        raise ValueError(f"Arkusz {sheet.title!r}: zduplikowane godziny.")
    return hours


def _read_clients(
    sheet: object, allowed_client_ids: Optional[AbstractSet[str]] = None,
    *, profiles_in_sheet: bool = False,
) -> Tuple[Tuple[str, ...], frozenset[str]]:
    clients: List[str] = []
    all_clients = set()
    first_row, column, letter = (3, 2, "B") if profiles_in_sheet else (2, 1, "A")
    for row_number, (raw_client,) in enumerate(
        sheet.iter_rows(min_row=first_row, min_col=column, max_col=column, values_only=True), start=first_row
    ):
        client_id = _client_identifier(raw_client, f"Klienci!{letter}{row_number}")
        if not client_id:
            continue
        if client_id in all_clients:
            if profiles_in_sheet:
                # Ten sam PPE pojawia się dla kolejnych dat. Duplikaty całego
                # klucza (OSD, PPE, data) kontroluje importer profili.
                continue
            raise ValueError(f"Arkusz Klienci zawiera duplikat {client_id!r}.")
        all_clients.add(client_id)
        if allowed_client_ids is None or client_id in allowed_client_ids:
            clients.append(client_id)
    if not clients:
        if allowed_client_ids is None:
            raise ValueError(f"Arkusz Klienci nie zawiera żadnego klienta od {letter}{first_row} w dół.")
        raise ValueError(
            "Po połączeniu z aktualną listą zgód w arkuszu Klienci nie został "
            "żaden dopuszczony PPE."
        )
    return tuple(clients), frozenset(all_clients)


def _read_targets(
    sheet: object,
    requested_branch: Optional[str],
) -> Tuple[Tuple[str, ...], str, Tuple[Period, ...], np.ndarray]:
    """Read the target while preserving GAMS' global date/hour domain."""

    hours = _hours_from_header(sheet, header_row=2, first_hour_column=3)
    normalized_request = (
        _identifier(requested_branch) if requested_branch is not None else None
    )
    branches: List[str] = []
    seen_branches = set()
    dates: "OrderedDict[object, None]" = OrderedDict()
    seen_branch_dates = set()
    target_rows: "OrderedDict[object, Tuple[float, ...]]" = OrderedDict()
    tentative_branch: Optional[str] = normalized_request
    max_column = 2 + len(hours)
    for row_number, row in enumerate(
        sheet.iter_rows(min_row=3, max_col=max_column, values_only=True), start=3
    ):
        branch = _identifier(row[0])
        history_date = _axis_key(row[1])
        if not branch and history_date in (None, ""):
            continue
        if not branch or history_date in (None, ""):
            raise ValueError(
                f"Arkusz Oddzial, wiersz {row_number}: brak oddziału albo daty."
            )
        if branch not in seen_branches:
            branches.append(branch)
            seen_branches.add(branch)
            if tentative_branch is None:
                tentative_branch = branch
        dates.setdefault(history_date, None)

        branch_date = (branch, history_date)
        if branch_date in seen_branch_dates:
            raise ValueError(
                "Arkusz Oddzial zawiera duplikat daty "
                f"{history_date!r} dla {branch!r}."
            )
        seen_branch_dates.add(branch_date)

        # GDXXRW reads all rows before the GAMS model chooses an OSD, so a bad
        # numeric value in another branch must not be silently ignored.
        values = tuple(
            _number(row[offset], f"Oddzial!{row_number}, godzina {hour}")
            for offset, hour in enumerate(hours, start=2)
        )
        if branch == tentative_branch:
            target_rows[history_date] = values

    if not branches:
        raise ValueError("Arkusz Oddzial nie zawiera danych od wiersza 3.")
    if normalized_request is None and len(branches) != 1:
        raise ValueError(
            "Plik zawiera kilka oddziałów; podaj --branch. Dostępne: "
            + ", ".join(branches)
        )
    selected_branch = normalized_request or branches[0]
    if selected_branch not in seen_branches:
        raise ValueError(
            f"Nie znaleziono oddziału {selected_branch!r}. Dostępne: "
            + ", ".join(branches)
        )

    # s_Data_Hist is a global GAMS set loaded from the complete column B.
    # Therefore even dates present only for another OSD generate equations,
    # with a zero target for the selected OSD.
    periods = tuple((history_date, hour) for history_date in dates for hour in hours)
    target = np.zeros(len(periods), dtype=float)
    hour_count = len(hours)
    date_positions = {history_date: index for index, history_date in enumerate(dates)}
    for history_date, values in target_rows.items():
        start = date_positions[history_date] * hour_count
        target[start : start + hour_count] = values
    return tuple(branches), selected_branch, periods, target


def _read_restrictions(
    sheet: object,
) -> Tuple[frozenset[str], Dict[str, float]]:
    excluded = set()
    required: Dict[str, float] = {}
    for row_number, row in enumerate(
        sheet.iter_rows(min_row=2, min_col=1, max_col=6, values_only=True), start=2
    ):
        excluded_id = _client_identifier(row[0], f"Ograniczenia!A{row_number}")
        if excluded_id:
            excluded.add(excluded_id)

        required_id = _client_identifier(row[4], f"Ograniczenia!E{row_number}")
        raw_weight = row[5]
        if required_id:
            if required_id in required:
                raise ValueError(
                    f"Arkusz Ograniczenia zawiera duplikat wymaganego klienta "
                    f"{required_id!r}."
                )
            if raw_weight is None or (isinstance(raw_weight, str) and not raw_weight.strip()):
                raise ValueError(
                    f"Ograniczenia!E{row_number}: wymagany klient nie ma wagi w kolumnie F."
                )
            required[required_id] = _number(
                raw_weight, f"Ograniczenia!F{row_number}"
            )
        elif raw_weight not in (None, ""):
            raise ValueError(
                f"Ograniczenia!F{row_number}: waga bez identyfikatora w kolumnie E."
            )
    return frozenset(excluded), required


def _add_profile_sheet(
    sheet: object,
    branch: str,
    known_branches: frozenset[str],
    known_client_ids: frozenset[str],
    client_positions: Dict[str, int],
    period_positions: Dict[Period, int],
    profiles: np.ndarray,
    on_stage: Optional[Callable[[str], None]] = None,
) -> None:
    hours = _hours_from_header(sheet, header_row=2, first_hour_column=4)
    max_column = 3 + len(hours)
    seen_rows = set()
    for row_number, row in enumerate(
        sheet.iter_rows(min_row=3, max_col=max_column, values_only=True), start=3
    ):
        if on_stage is not None and row_number % 100_000 == 0:
            on_stage(
                f"{sheet.title}: przetworzono {row_number - 2:,} wierszy..."
            )
        row_branch = _identifier(row[0])
        client_id = _client_identifier(row[1], f"{sheet.title}!B{row_number}")
        history_date = _axis_key(row[2])
        if not row_branch and not client_id and history_date in (None, ""):
            continue
        if not row_branch or not client_id or history_date in (None, ""):
            raise ValueError(
                f"{sheet.title}, wiersz {row_number}: brak oddziału, klienta albo daty."
            )
        if client_id not in known_client_ids:
            raise ValueError(
                f"{sheet.title}!B{row_number}: klient {client_id!r} nie występuje "
                "w arkuszu Klienci."
            )
        if row_branch not in known_branches:
            raise ValueError(
                f"{sheet.title}!A{row_number}: oddział {row_branch!r} nie występuje "
                "w arkuszu Oddzial."
            )
        row_key = (row_branch, client_id, history_date)
        if row_key in seen_rows:
            raise ValueError(
                f"{sheet.title} zawiera duplikat klienta {client_id!r} dla daty "
                f"{history_date!r} i oddziału {row_branch!r}."
            )
        seen_rows.add(row_key)
        values = tuple(
            _number(row[offset], f"{sheet.title}!{row_number}, godzina {hour}")
            for offset, hour in enumerate(hours, start=3)
        )
        for hour, value in zip(hours, values):
            period = (history_date, hour)
            if period not in period_positions and value != 0:
                raise ValueError(
                    f"{sheet.title}, wiersz {row_number}: okres {period!r} "
                    "nie występuje w bilansie oddziałów."
                )
        if row_branch != branch or client_id not in client_positions:
            continue

        client_index = client_positions[client_id]
        for hour, value in zip(hours, values):
            period = (history_date, hour)
            if period not in period_positions:
                continue
            profiles[period_positions[period], client_index] += value


_CELL_REFERENCE = re.compile(r"^([A-Z]+)([1-9][0-9]*)$")


def _cell_position(reference: str) -> Tuple[int, int]:
    """Return a one-based ``(row, column)`` pair for an A1 reference."""

    match = _CELL_REFERENCE.match(reference.upper())
    if not match:
        return 0, 0
    column = 0
    for character in match.group(1):
        column = column * 26 + ord(character) - ord("A") + 1
    return int(match.group(2)), column


def _uncached_formula_cells(
    source: Path,
    ranges: Dict[str, Tuple[int, int, int]],
    maximum_examples: int = 10,
) -> Tuple[int, Tuple[str, ...]]:
    """Find relevant formulas without a cached value, using bounded memory.

    In ``data_only`` mode openpyxl exposes both a truly empty cell and an
    unevaluated formula as ``None``.  Blank data cells intentionally mean zero
    in the legacy GDXXRW import, whereas a formula without a cached result is
    unsafe.  The XLSX XML retains this distinction, so inspect it directly.
    """

    try:
        with ZipFile(source) as archive:
            with archive.open("xl/workbook.xml") as stream:
                workbook_root = ElementTree.parse(stream).getroot()
            with archive.open("xl/_rels/workbook.xml.rels") as stream:
                relationship_root = ElementTree.parse(stream).getroot()

            targets = {
                relationship.attrib["Id"]: relationship.attrib["Target"]
                for relationship in relationship_root
            }
            relationship_id = (
                "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"
            )
            sheet_paths: Dict[str, str] = {}
            for element in workbook_root.iter():
                if element.tag.rsplit("}", 1)[-1] != "sheet":
                    continue
                name = element.attrib.get("name")
                relation = element.attrib.get(relationship_id)
                if name not in ranges or relation not in targets:
                    continue
                target = targets[relation]
                sheet_paths[name] = (
                    target.lstrip("/")
                    if target.startswith("/")
                    else posix_normpath(posix_join("xl", target))
                )

            count = 0
            examples: List[str] = []
            for sheet_name, bounds in ranges.items():
                sheet_path = sheet_paths.get(sheet_name)
                if not sheet_path:
                    continue
                minimum_row, minimum_column, maximum_column = bounds
                with archive.open(sheet_path) as worksheet_xml:
                    for _, cell in ElementTree.iterparse(
                        worksheet_xml, events=("end",)
                    ):
                        local_name = cell.tag.rsplit("}", 1)[-1]
                        if local_name == "row":
                            cell.clear()
                            continue
                        if local_name != "c":
                            continue
                        reference = cell.attrib.get("r", "")
                        row_number, column_number = _cell_position(reference)
                        if not (
                            row_number >= minimum_row
                            and minimum_column <= column_number <= maximum_column
                        ):
                            cell.clear()
                            continue

                        formula = None
                        cached_value = None
                        for child in cell:
                            child_name = child.tag.rsplit("}", 1)[-1]
                            if child_name == "f":
                                formula = child
                            elif child_name == "v":
                                cached_value = child
                        if formula is not None and (
                            cached_value is None or cached_value.text is None
                        ):
                            count += 1
                            if len(examples) < maximum_examples:
                                examples.append(f"{sheet_name}!{reference}")
                        cell.clear()
            return count, tuple(examples)
    except (BadZipFile, KeyError, ElementTree.ParseError) as exc:
        raise ValueError(f"Nieprawidłowy albo uszkodzony plik XLSX: {source}.") from exc


def load_gams_workbook(
    path: str | Path,
    branch: Optional[str] = None,
    on_stage: Optional[Callable[[str], None]] = None,
    allowed_client_ids: Optional[AbstractSet[str]] = None,
) -> SelectionData:
    """Czyta układ dzielony v8 lub trzyarkuszowy wariant z profilami w Klienci.

    Odtwarzane zakresy to: ``Oddzial!A2``, ``Klienci!A2``,
    ``Klienci_part_1!A2``, ``Klienci_part_2!A2`` i ``Ograniczenia!A2/F2``.
    Okresy są globalnym iloczynem wszystkich dat i godzin, tak jak domeny
    ``s_Data_Hist`` i ``s_h`` w GAMS. Puste komórki liczbowe oznaczają zero;
    formuła bez wyniku zapisanego w pliku jest odrzucana. Plik i kontrola
    formuł są czytane strumieniowo, a w pamięci zachowywany jest target tylko
    wybranego oddziału. Końcowa macierz profili pozostaje jednak gęsta.
    """

    try:
        from openpyxl import load_workbook
        from openpyxl.utils.exceptions import InvalidFileException
    except ImportError as exc:  # pragma: no cover - zależy od środowiska
        raise RuntimeError("Do odczytu XLSX wymagany jest openpyxl.") from exc

    source = Path(path)
    if not source.exists():
        raise FileNotFoundError(source)
    if on_stage is not None:
        on_stage("otwieranie skoroszytu...")
    try:
        workbook = load_workbook(source, read_only=True, data_only=True)
    except (BadZipFile, InvalidFileException, OSError, KeyError, ValueError) as exc:
        raise ValueError(
            f"Nie można otworzyć pliku {source} jako poprawnego skoroszytu XLSX."
        ) from exc
    try:
        sheets = _resolve_sheet_names(workbook)
    except ValueError as exc:
        workbook.close()
        raise ValueError(f"{source.name}: {exc}") from exc

    try:
        profile_sheets = _profile_sheet_names(sheets)
        inline = "Klienci_part_1" not in sheets
        if on_stage is not None:
            on_stage("układ: profile w Klienci (PPE w B, data w C, godziny od D)" if inline else "układ: Klienci_part_1 + Klienci_part_2")
            on_stage("sprawdzanie formuł i pamięci podręcznej Excela...")
        formula_count, formula_examples = _uncached_formula_cells(
            source,
            {
                sheets["Oddzial"]: (
                    2,
                    1,
                    workbook[sheets["Oddzial"]].max_column,
                ),
                sheets["Klienci"]: (2, 1, 1),
                **{name: (2, 1, workbook[name].max_column) for name in profile_sheets},
                sheets["Ograniczenia"]: (2, 1, 6),
            },
        )
        if formula_count:
            examples = ", ".join(formula_examples)
            remaining = formula_count - len(formula_examples)
            if remaining:
                examples += f" oraz {remaining} kolejnych"
            raise ValueError(
                f"Plik zawiera formuły bez zapisanej wartości ({formula_count}): "
                f"{examples}. Otwórz plik w Excelu, przelicz i zapisz go "
                "przed uruchomieniem."
            )

        if on_stage is not None:
            on_stage("wczytywanie klientów i bilansu oddziału...")
        clients, all_client_ids = _read_clients(
            workbook[sheets["Klienci"]], allowed_client_ids, profiles_in_sheet=inline,
        )
        (
            branches,
            selected_branch,
            periods,
            branch_balance,
        ) = _read_targets(workbook[sheets["Oddzial"]], branch)
        try:
            profiles = np.zeros((len(periods), len(clients)), dtype=float)
        except (MemoryError, ValueError) as exc:
            gibibytes = len(periods) * len(clients) * 8 / 1024**3
            raise MemoryError(
                "Nie można utworzyć gęstej macierzy profili "
                f"{len(periods)} x {len(clients)} (około {gibibytes:.2f} GiB)."
            ) from exc
        client_positions = {client_id: index for index, client_id in enumerate(clients)}
        period_positions = {period: index for index, period in enumerate(periods)}
        for name in profile_sheets:
            if on_stage is not None:
                on_stage(f"wczytywanie {name}...")
            _add_profile_sheet(
                workbook[name], selected_branch, frozenset(branches),
                all_client_ids, client_positions, period_positions, profiles, on_stage,
            )
        if on_stage is not None:
            on_stage("wczytywanie ograniczeń...")
        excluded_all, required_all = _read_restrictions(
            workbook[sheets["Ograniczenia"]]
        )
        selected_ids = set(clients)
        excluded = frozenset(excluded_all & selected_ids)
        required_without_consent = (set(required_all) & set(all_client_ids)) - selected_ids
        if allowed_client_ids is not None and required_without_consent:
            examples = ", ".join(sorted(required_without_consent)[:10])
            raise ValueError(
                "Lista klientów wymaganych zawiera PPE bez aktualnej zgody albo "
                f"spoza wybranego OSD: {examples}."
            )
        required = {
            client_id: weight
            for client_id, weight in required_all.items()
            if client_id in selected_ids
        }

    except ValueError as exc:
        raise ValueError(f"{source.name}: {exc}") from exc
    finally:
        workbook.close()

    if on_stage is not None:
        on_stage("końcowa walidacja domen i wartości...")
    data = SelectionData(
        branch=selected_branch,
        client_ids=clients,
        periods=periods,
        client_balance=profiles,
        branch_balance=branch_balance,
        excluded=excluded,
        required_weights=required,
    )
    data.validate()
    return data
