from __future__ import annotations

import os
import tempfile
from dataclasses import dataclass
from pathlib import Path
from typing import Mapping, Sequence, Tuple

import pandas as pd

from .consent import BRANCH_LABELS, canonical_branch
from .model import SelectionResult
from .replacement import ReplacementResult
from .report import balance_frame, diagnostics_frame, selected_frame


EXCEL_MAX_ROWS = 1_048_576


@dataclass(frozen=True)
class BranchOutput:
    selection: SelectionResult
    replacements: ReplacementResult
    source_files: Tuple[str, ...]
    source_clients: int
    consent_clients: int
    legacy_excluded: int
    eligible_clients: int
    target_clients: int
    quality: Mapping[str, pd.DataFrame] | None = None
    with_reserves: bool = True


def _plan_status(output: BranchOutput | None, row: dict) -> str:
    if output is None:
        return "Brak kandydatów — bilans nie jest oceniony" if row["eligible"] == 0 else "Poza tym raportem"
    if output.selection.selected_count != row["target"]:
        return "Niekompletny"
    if not output.with_reserves:
        return "Kompletny wybór bez zastępców"
    return "Komplet par" if len(output.replacements.matches) == row["target"] else "Niekompletny"


def _frames(outputs: Sequence[BranchOutput]) -> Mapping[str, pd.DataFrame]:
    selected_parts = [selected_frame(output.selection) for output in outputs]
    balance_parts = [output.quality["Bilans"] if output.quality else balance_frame(output.selection) for output in outputs]
    solver_parts = [diagnostics_frame(output.selection) for output in outputs]
    replacement_rows = []
    replacement_diagnostics = []
    funnel_rows = []
    summary_rows = []
    for output in outputs:
        key = canonical_branch(output.selection.branch)
        label = BRANCH_LABELS[key]
        for match in output.replacements.matches:
            replacement_rows.append(
                {
                    "OSD": label,
                    "grupa": match.group,
                    "klient_podstawowy": match.primary_id,
                    "waga_podstawowa": match.primary_weight,
                    "klient_zastepczy": match.reserve_id,
                    "waga_zastepcza": match.reserve_weight,
                    "blad_L1_podmiany": match.l1_error,
                    "MAE_podmiany": match.mae,
                    "blad_L1_wzgledny": match.normalized_l1,
                    "korelacja": match.correlation,
                }
            )
        for diagnostic in output.replacements.diagnostics:
            replacement_diagnostics.append(
                {
                    "OSD": label,
                    "grupa": diagnostic.group,
                    "L1_model_podstawowy": diagnostic.l1_primary,
                    "L1_po_podmianie_wszystkich": diagnostic.l1_after_replacement,
                    "MAE_model_podstawowy": diagnostic.mae_primary,
                    "MAE_po_podmianie_wszystkich": diagnostic.mae_after_replacement,
                    "max_blad_model_podstawowy": diagnostic.max_error_primary,
                    "max_blad_po_podmianie": diagnostic.max_error_after_replacement,
                }
            )
        funnel_rows.append(
            {
                "OSD": label,
                "pliki_zrodlowe": "; ".join(output.source_files),
                "klienci_w_modelu_zrodlowym": output.source_clients,
                "PPE_zgodne_z_plikiem_zgod": output.consent_clients,
                "wykluczeni_legacy": output.legacy_excluded,
                "kandydaci_po_filtrach": output.eligible_clients,
                "wybrani": output.selection.selected_count,
                "zastepcy": len(output.replacements.matches),
                "cel_dla_OSD": output.target_clients,
            }
        )
        group_diags = output.replacements.diagnostics
        summary_rows.append(
            {
                "OSD": label,
                "cel": output.target_clients,
                "wybrani": output.selection.selected_count,
                "zastepcy": len(output.replacements.matches),
                "liczba_grup": len(output.selection.groups),
                "wszystkie_rozwiazania_optymalne": all(
                    group.optimal for group in output.selection.groups
                ),
                "srednie_MAE_model": (
                    sum(item.mae_primary for item in group_diags) / len(group_diags)
                    if group_diags
                    else (float(output.quality["Reprezentatywnosc"].loc[
                        output.quality["Reprezentatywnosc"]["zakres"] == "Dopasowanie", "MAE"
                    ].mean()) if output.quality else None)
                ),
                "srednie_MAE_po_podmianie": (
                    sum(item.mae_after_replacement for item in group_diags)
                    / len(group_diags)
                    if group_diags
                    else None
                ),
                "kompletny_wybor_OSD": output.selection.selected_count == output.target_clients,
                "komplet_zastepcow": len(output.replacements.matches) == output.target_clients if output.with_reserves else None,
                "tryb_zastepcow": "Włączony" if output.with_reserves else "Wyłączony — ocena dostępnej puli",
                "ocena_testowa": "Test na odłożonych dniach" if output.quality and
                    "Test poza doborem" in output.quality["Reprezentatywnosc"]["zakres"].values else "Brak testu poza doborem",
            }
        )
    frames = {
        "Podsumowanie": pd.DataFrame(summary_rows),
        "Wybrani": pd.concat(selected_parts, ignore_index=True),
        "Zastepcy": pd.DataFrame(replacement_rows),
        "Bilans": pd.concat(balance_parts, ignore_index=True),
        "Diagnostyka": pd.DataFrame(replacement_diagnostics),
        "Solver": pd.concat(solver_parts, ignore_index=True),
        "Lejek_danych": pd.DataFrame(funnel_rows),
    }
    for name in ("Reprezentatywnosc", "Jakosc_dni_godzin", "Podmiany_pojedyncze"):
        parts = [output.quality[name] for output in outputs if output.quality and not output.quality[name].empty]
        if parts:
            frames[name] = pd.concat(parts, ignore_index=True)
    return frames


def _mail_text(outputs: Sequence[BranchOutput], total_target: int) -> str:
    selected = sum(output.selection.selected_count for output in outputs)
    reserves = sum(len(output.replacements.matches) for output in outputs)
    optimal = sum(
        group.optimal for output in outputs for group in output.selection.groups
    )
    group_count = sum(len(output.selection.groups) for output in outputs)
    distribution = ", ".join(
        f"{BRANCH_LABELS[canonical_branch(output.selection.branch)]}: "
        f"{output.selection.selected_count}"
        for output in outputs
    )
    reserve_text = "Dobór zastępców wyłączony — ocena dostępnej puli PPE. "
    if any(output.with_reserves for output in outputs):
        reserve_text = (
            ("Dobór zastępców jest kompletny dla zapisanych klientów. " if selected == reserves else
             "Dobór zastępców NIE JEST kompletny; ten zapis nie stanowi gotowej listy par. ")
            + "Zastępcy pochodzą wyłącznie spośród PPE z aktualną zgodą. "
            "Zastępcy zostali przypisani wspólnie w obrębie OSD z krótkich list podobnych "
            "profili, minimalizując sumę względnych błędów L1 podmian na tych listach. "
            "Nie dowodzono optymalności względem wszystkich możliwych par. "
            "Błąd podmiany wszystkich klientów jest scenariuszem innym niż podmiana jednej osoby. "
        )
    return (
        "Model klientów reprezentatywnych został uruchomiony po ograniczeniu zbioru "
        "do PPE znajdujących się w aktualnym pliku zgód marketingowych i po "
        "zastosowaniu technicznych wykluczeń z arkuszy wejściowych. "
        f"Docelowa liczba klientów wynosiła {total_target}; wybrano dokładnie "
        f"{selected}. Podział na OSD: {distribution}.\n\n"
        "Dla każdego OSD kandydaci zostali rozdzieleni round-robin na "
        "niezależne grupy zgodnie z konfiguracją. W każdej grupie minimalizowano sumę bezwzględnych "
        "różnic godzinowych pomiędzy historycznym bilansem OSD a ważoną sumą "
        "profili klientów, z zachowaniem limitów wag i udziału pojedynczego PPE. "
        f"Solver zakończył w zadanej tolerancji luki MIP w {optimal} z {group_count} grup "
        "(szczegółowe luki i statusy w arkuszu Solver).\n\n"
        f"Liczba zapisanych klientów podstawowych: {selected}; liczba różnych zastępców: {reserves}. "
        + reserve_text +
        "Szczegółowe wyniki, lejek danych, statusy "
        "solvera i wpływ podmian znajdują się w załączonym skoroszycie. "
        "Każda grupa osobno odtwarza cały bilans oddziału; wag różnych grup nie należy sumować. "
        "Ocena na danych użytych do doboru jest dopasowaniem historycznym. "
        "Tylko wiersze oznaczone jako test poza doborem dotyczą dni niewykorzystanych przez solver ani dobór zastępców."
    )


def _style_workbook(workbook: object) -> None:
    from openpyxl.styles import Alignment, Font, PatternFill

    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    for sheet in workbook.worksheets:
        sheet.freeze_panes = "A2"
        if sheet.max_row >= 1 and sheet.max_column >= 1:
            sheet.auto_filter.ref = sheet.dimensions
            for cell in sheet[1]:
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(wrap_text=True, vertical="top")
            sheet.row_dimensions[1].height = 46
        for column_cells in sheet.iter_cols(
            min_row=1, max_row=min(sheet.max_row, 200), max_col=sheet.max_column
        ):
            width = min(
                45,
                max(10, max(len(str(cell.value or "")) for cell in column_cells) + 2),
            )
            sheet.column_dimensions[column_cells[0].column_letter].width = width
        headers = [cell.value for cell in sheet[1]]
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                header = str(headers[cell.column - 1] or "")
                if isinstance(cell.value, float):
                    cell.number_format = "0.0%" if ("wzgledny" in header or header == "mip_gap") else "#,##0.00"
                elif isinstance(cell.value, int) and not isinstance(cell.value, bool):
                    cell.number_format = "#,##0"
                if cell.data_type == "f":
                    # Identyfikatory i tekst ze źródeł mają pozostać tekstem, nigdy formułami.
                    cell.data_type = "s"


def _add_charts(workbook) -> None:
    from openpyxl.chart import LineChart, Reference
    source = workbook["Bilans"]
    if source.max_column < 10 or source.max_row < 2:
        return
    sheet = workbook.create_sheet("Wykresy")
    sheet["A1"] = "Historia i model; każda grupa osobno odtwarza cały oddział. Wykres pokazuje do 1000 pierwszych okresów grupy; metryki obejmują wszystkie."
    sheet["A1"].alignment = __import__("openpyxl").styles.Alignment(wrap_text=True)
    sheet.column_dimensions["A"].width = 120
    sheet.row_dimensions[1].height = 34
    starts = []
    previous = None
    for row in range(2, source.max_row + 1):
        key = (source.cell(row, 1).value, source.cell(row, 2).value)
        if key != previous:
            starts.append((row, key))
            previous = key
    for index, (start, key) in enumerate(starts):
        end = min(start + 999, starts[index + 1][0] - 1 if index + 1 < len(starts) else source.max_row)
        chart = LineChart()
        chart.title = f"{key[0]} — {key[1]}"
        chart.y_axis.title = "Bilans w jednostce danych źródłowych"
        chart.x_axis.title = "Data / godzina"
        chart.height, chart.width = 9, 25
        chart.legend.position = "b"
        for column, title in ((5, "Historia"), (6, "Model")):
            chart.add_data(Reference(source, min_col=column, min_row=start, max_row=end))
            from openpyxl.chart.series import SeriesLabel
            chart.series[-1].tx = SeriesLabel(v=title)
        chart.set_categories(Reference(source, min_col=10, min_row=start, max_row=end))
        chart.x_axis.tickLblSkip = max(1, (end - start + 1) // 8)
        sheet.add_chart(chart, f"A{3 + index * 19}")


def write_portfolio_workbook(
    outputs: Sequence[BranchOutput],
    path: str | Path,
    total_target: int,
    *,
    overwrite: bool = False,
    partial: bool = False,
    run_metadata: Mapping[str, object] | None = None,
    allocation_plan: dict | None = None,
) -> Path:
    if not outputs:
        raise ValueError("Brak wyników do zapisania.")
    target = Path(path)
    if target.exists() and not overwrite:
        raise FileExistsError(f"Plik {target} już istnieje.")
    target.parent.mkdir(parents=True, exist_ok=True)
    frames = _frames(outputs)
    run_metadata = dict(run_metadata or {})
    coverage = run_metadata.pop("pokrycie_OSD", None)
    if coverage:
        frames["Pokrycie_OSD"] = pd.DataFrame([{"OSD": key, **row} for key, row in coverage.items()])
    if allocation_plan:
        current = {canonical_branch(item.selection.branch): item for item in outputs}
        frames["Plan_podzialu"] = pd.DataFrame([
            {"OSD": BRANCH_LABELS[key], "klienci_w_planie": row["target"],
             "zastepcy_w_planie": row["target"] if allocation_plan.get("with_reserves", True) else 0,
             "kandydaci_przy_utworzeniu": row["eligible"],
             "zapisani_w_tym_raporcie": current[key].selection.selected_count if key in current else 0,
             "status_w_tym_raporcie": _plan_status(current.get(key), row),
             "plan_id": allocation_plan["plan_id"]}
            for key, row in allocation_plan["branches"].items()
        ])
    oversized = {
        name: len(frame) + 1
        for name, frame in frames.items()
        if len(frame) + 1 > EXCEL_MAX_ROWS
    }
    if oversized:
        details = ", ".join(f"{name}: {rows}" for name, rows in oversized.items())
        raise ValueError(
            f"Raport przekroczyłby limit {EXCEL_MAX_ROWS} wierszy Excela ({details})."
        )
    descriptor, temporary_name = tempfile.mkstemp(
        prefix=f".{target.stem}.", suffix=".tmp.xlsx", dir=target.parent
    )
    os.close(descriptor)
    temporary = Path(temporary_name)
    try:
        with pd.ExcelWriter(temporary, engine="openpyxl") as writer:
            for sheet_name, frame in frames.items():
                frame.to_excel(writer, sheet_name=sheet_name, index=False)
            pd.DataFrame(
                {
                    "parametr": [
                        "docelowa_liczba_klientow",
                        "liczba_OSD",
                        "zasada_zgod",
                        "zasada_zastepcow",
                        "zakres_raportu",
                        *(run_metadata or {}).keys(),
                    ],
                    "wartosc": [
                        total_target,
                        len(outputs),
                        "Wyłącznie aktualna lista PPE; korekta ma pierwszeństwo; pierwotny eksport służy tylko do mapowania OSD",
                        "Jeden unikalny zastępca na klienta, w tym samym OSD" if any(item.with_reserves for item in outputs) else "Dobór zastępców wyłączony",
                        "CZĘŚCIOWY — sprawdź kompletność wyboru i zastępców w Podsumowaniu" if partial else
                        ("Nie oceniono OSD bez kandydatów — sprawdź Pokrycie_OSD" if run_metadata.get("OSD_bez_kandydatow") else "Pełny przebieg wskazanych OSD"),
                        *(run_metadata or {}).values(),
                    ],
                }
            ).to_excel(writer, sheet_name="Konfiguracja", index=False)
            mail_text = _mail_text(outputs, total_target)
            if run_metadata.get("OSD_bez_kandydatow"):
                mail_text = ("Nie oceniono pełnego portfela. OSD bez kandydatów: "
                             + str(run_metadata["OSD_bez_kandydatow"]) + ".\n\n" + mail_text)
            if partial:
                mail_text = "RAPORT CZĘŚCIOWY: obejmuje zapisane grupy/OSD. Kompletność wyboru i zastępców podano w Podsumowaniu.\n\n" + mail_text
            pd.DataFrame(
                {"Tekst do maila (bez identyfikatorów PPE)": mail_text.split("\n")}
            ).to_excel(writer, sheet_name="Tekst_do_maila", index=False)
            _style_workbook(writer.book)
            _add_charts(writer.book)
        if overwrite:
            os.replace(temporary, target)
        else:
            os.link(temporary, target)
    finally:
        temporary.unlink(missing_ok=True)
    return target
