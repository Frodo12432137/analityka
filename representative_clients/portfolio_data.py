from __future__ import annotations

from collections import OrderedDict
from pathlib import Path
from typing import Dict, Iterable, Mapping, Sequence, Tuple

import numpy as np

from .consent import BRANCH_ORDER, canonical_branch
from .grouping import round_robin_groups
from .model import SelectionData


def inspect_workbook_clients(
    path: str | Path, allowed_client_ids: Iterable[str]
) -> Tuple[frozenset[str], frozenset[str], frozenset[str], frozenset[str]]:
    """Czyta tylko domenę klientów i ograniczenia, bez profili godzinowych."""

    from openpyxl import load_workbook

    from .legacy_excel import _read_clients, _read_restrictions, _resolve_sheet_names

    workbook = load_workbook(Path(path), read_only=True, data_only=True)
    try:
        sheets = _resolve_sheet_names(workbook)
        _, all_clients = _read_clients(
            workbook[sheets["Klienci"]], profiles_in_sheet="Klienci_part_1" not in sheets,
        )
        excluded, required = _read_restrictions(workbook[sheets["Ograniczenia"]])
        client_set = all_clients & frozenset(allowed_client_ids)
        missing_consent = (set(required) & all_clients) - client_set
        if missing_consent:
            raise ValueError(
                f"{len(missing_consent)} wymaganych PPE nie ma "
                "w pliku zgód dla tego OSD. Sprawdź Ograniczenia!E:F."
            )
        if set(required) & excluded:
            raise ValueError("PPE jednocześnie wymagany i wykluczony.")
        return (
            all_clients,
            client_set,
            frozenset(excluded & client_set),
            frozenset(set(required) & client_set),
        )
    except ValueError as exc:
        raise ValueError(f"{Path(path).name}: {exc}") from exc
    finally:
        workbook.close()


def discover_workbook_branches(path: str | Path) -> Tuple[str, ...]:
    """Zwraca nazwy oddziałów z arkusza Oddzial bez wczytywania profili."""

    from openpyxl import load_workbook

    source = Path(path)
    workbook = load_workbook(source, read_only=True, data_only=True)
    try:
        sheet_name = next(
            (name for name in ("Oddzial", "Oddział") if name in workbook.sheetnames),
            None,
        )
        if sheet_name is None:
            raise ValueError(f"{source.name}: brak arkusza Oddzial/Oddział.")
        values: "OrderedDict[str, None]" = OrderedDict()
        for (raw_branch,) in workbook[sheet_name].iter_rows(
            min_row=3, min_col=1, max_col=1, values_only=True
        ):
            if raw_branch not in (None, ""):
                values.setdefault(str(raw_branch).strip(), None)
    finally:
        workbook.close()
    if not values:
        raise ValueError(f"{source.name}: brak nazw OSD w arkuszu Oddzial.")
    return tuple(values)


def merge_selection_data(parts: Sequence[SelectionData]) -> SelectionData:
    """Łączy części jednego OSD, odrzucając sprzeczne duplikaty."""

    if not parts:
        raise ValueError("Brak części do połączenia.")
    for part in parts:
        part.validate()
    if len(parts) == 1:
        return parts[0]
    branch_key = canonical_branch(parts[0].branch)
    periods = parts[0].periods
    target = np.asarray(parts[0].branch_balance, dtype=float)
    clients: list[str] = []
    profiles: list[np.ndarray] = []
    positions: Dict[str, int] = {}
    excluded: set[str] = set()
    required: Dict[str, float] = {}
    for part in parts:
        if canonical_branch(part.branch) != branch_key:
            raise ValueError("Nie można połączyć plików należących do różnych OSD.")
        if part.periods != periods:
            raise ValueError(
                f"Części OSD {part.branch} mają różne daty lub godziny."
            )
        if not np.allclose(part.branch_balance, target, rtol=1e-9, atol=1e-9):
            raise ValueError(f"Części OSD {part.branch} mają różny bilans oddziału.")
        for index, client_id in enumerate(part.client_ids):
            profile = np.asarray(part.client_balance[:, index], dtype=float)
            previous = positions.get(client_id)
            if previous is None:
                positions[client_id] = len(clients)
                clients.append(client_id)
                profiles.append(profile.copy())
            elif not np.allclose(profiles[previous], profile, rtol=1e-9, atol=1e-9):
                raise ValueError(
                    f"PPE {client_id!r} ma różne profile w częściach OSD {part.branch}."
                )
        excluded.update(part.excluded)
        for client_id, weight in part.required_weights.items():
            if client_id in required and not np.isclose(required[client_id], weight):
                raise ValueError(
                    f"PPE {client_id!r} ma dwie różne wagi wymagane."
                )
            required[client_id] = float(weight)
    merged = SelectionData(
        branch=parts[0].branch,
        client_ids=tuple(clients),
        periods=periods,
        client_balance=np.column_stack(profiles),
        branch_balance=target.copy(),
        excluded=frozenset(excluded),
        required_weights=required,
    )
    merged.validate()
    return merged


def allocate_bounded(
    total: int,
    weights: Mapping[str, int],
    lower: Mapping[str, int],
    upper: Mapping[str, int],
    order: Iterable[str],
) -> Dict[str, int]:
    """Deterministyczny podział proporcjonalny z dolnymi i górnymi limitami."""

    keys = tuple(key for key in order if key in weights)
    if not keys:
        raise ValueError("Brak pozycji do podziału limitu.")
    allocation = {key: int(lower.get(key, 0)) for key in keys}
    for key in keys:
        if allocation[key] < 0 or allocation[key] > int(upper[key]):
            raise ValueError(f"Sprzeczne limity dla {key}.")
    if total < sum(allocation.values()) or total > sum(int(upper[key]) for key in keys):
        raise ValueError(
            f"Nie da się rozdzielić dokładnie {total} klientów przy podanych pojemnościach."
        )
    remaining = total - sum(allocation.values())
    while remaining:
        open_keys = [key for key in keys if allocation[key] < int(upper[key])]
        denominator = sum(max(0, int(weights[key])) for key in open_keys)
        if denominator <= 0:
            denominator = len(open_keys)
            normalized = {key: 1 for key in open_keys}
        else:
            normalized = {key: max(0, int(weights[key])) for key in open_keys}
        ideals = {
            key: remaining * normalized[key] / denominator for key in open_keys
        }
        additions = {
            key: min(
                int(upper[key]) - allocation[key], int(np.floor(ideals[key]))
            )
            for key in open_keys
        }
        added = sum(additions.values())
        if added:
            for key, value in additions.items():
                allocation[key] += value
            remaining -= added
            continue
        ranked = sorted(
            open_keys,
            key=lambda key: (-(ideals[key] - np.floor(ideals[key])), keys.index(key)),
        )
        for key in ranked:
            if remaining == 0:
                break
            allocation[key] += 1
            remaining -= 1
    return allocation


def allocate_branch_targets(
    eligible_counts: Mapping[str, int], total: int
) -> Dict[str, int]:
    keys = tuple(key for key in BRANCH_ORDER if key in eligible_counts)
    if not keys:
        raise ValueError("Nie znaleziono żadnego obsługiwanego OSD.")
    capacity = {key: int(eligible_counts[key]) // 2 for key in keys}
    too_small = [key for key in keys if capacity[key] < 1]
    if too_small:
        raise ValueError(
            "Za mało klientów ze zgodą, aby wybrać klienta i unikalnego zastępcę: "
            + ", ".join(too_small)
        )
    return allocate_bounded(
        total,
        eligible_counts,
        {key: 1 for key in keys},
        capacity,
        keys,
    )


def allocate_group_quotas(data: SelectionData, total: int, group_count: int) -> Tuple[int, ...]:
    groups = round_robin_groups(data.client_ids, data.excluded, group_count)
    weights = {name: len(ids) for name, ids in groups.items()}
    required = {
        name: sum(client_id in data.required_weights for client_id in ids)
        for name, ids in groups.items()
    }
    allocation = allocate_bounded(
        total,
        weights,
        required,
        weights,
        groups.keys(),
    )
    return tuple(allocation[name] for name in groups)
