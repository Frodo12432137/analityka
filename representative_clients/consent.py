from __future__ import annotations

import re
import unicodedata
from dataclasses import dataclass
from pathlib import Path
from typing import Dict, Iterable, Mapping, Tuple


BRANCH_ORDER = (
    "bialystok",
    "lublin",
    "lodz_miasto",
    "lodz_teren",
    "rzeszow",
    "skarzysko_kamienna",
    "warszawa",
    "zamosc",
)

BRANCH_LABELS = {
    "bialystok": "Białystok",
    "lublin": "Lublin",
    "lodz_miasto": "Łódź Miasto",
    "lodz_teren": "Łódź Teren",
    "rzeszow": "Rzeszów",
    "skarzysko_kamienna": "Skarżysko-Kamienna",
    "warszawa": "Warszawa",
    "zamosc": "Zamość",
}


def _plain(value: object) -> str:
    text = str(value or "").translate(str.maketrans({"ł": "l", "Ł": "L"}))
    text = unicodedata.normalize("NFKD", text)
    text = "".join(character for character in text if not unicodedata.combining(character))
    return re.sub(r"[^a-z0-9]+", " ", text.casefold()).strip()


def _has_value(value: object) -> bool:
    return value is not None and str(value).strip() != ""


def canonical_branch(value: object) -> str:
    """Mapuje nazwy techniczne i opisowe OSD na osiem stałych kluczy."""

    text = _plain(value)
    compact = text.replace(" ", "")
    rules = (
        ("skarzysko_kamienna", ("skarzysko", "kamienna")),
        ("lodz_miasto", ("lodz", "miasto")),
        ("lodz_teren", ("lodz", "teren")),
        ("bialystok", ("bialystok",)),
        ("lublin", ("lublin",)),
        ("rzeszow", ("rzeszow",)),
        ("warszawa", ("warszawa",)),
        ("zamosc", ("zamosc",)),
    )
    for key, words in rules:
        if all(word in compact for word in words):
            return key
    raise ValueError(f"Nie rozpoznano OSD z wartości {value!r}.")


def technical_branch(key: str) -> str:
    mapping = {
        "bialystok": "PGED_Bialystok",
        "lublin": "PGED_Lublin",
        "lodz_miasto": "PGED_Lodz_Miasto",
        "lodz_teren": "PGED_Lodz_Teren",
        "rzeszow": "PGED_Rzeszow",
        "skarzysko_kamienna": "PGED_Skarzysko_Kamienna",
        "warszawa": "PGED_Warszawa",
        "zamosc": "PGED_Zamosc",
    }
    return mapping[key]


@dataclass(frozen=True)
class ConsentIndex:
    by_branch: Mapping[str, frozenset[str]]
    source_rows: int
    duplicate_rows: int
    corrected_rows: int = 0

    @property
    def unique_count(self) -> int:
        return sum(len(values) for values in self.by_branch.values())


def _header_positions(rows: Iterable[Tuple[object, ...]]) -> Tuple[int, int, int]:
    code_names = {"kod pp aktualny", "kod pp"}
    branch_names = {"osd nazwa pp aktualny", "osd nazwa", "osd"}
    for row_number, row in enumerate(rows, start=1):
        normalized = [_plain(value) for value in row]
        code_column = next(
            (index for index, value in enumerate(normalized) if value in code_names),
            None,
        )
        branch_column = next(
            (index for index, value in enumerate(normalized) if value in branch_names),
            None,
        )
        if code_column is not None and branch_column is not None:
            return row_number, code_column, branch_column
        if row_number >= 30:
            break
    raise ValueError(
        "Nie znaleziono nagłówków 'Kod PP Aktualny' i "
        "'OSD Nazwa - PP Aktualny' w pierwszych 30 wierszach pliku zgód."
    )


def _load_legacy_consent_workbook(path: str | Path) -> ConsentIndex:
    """Czyta skoroszyt zawierający wyłącznie aktualne zgody marketingowe."""

    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # pragma: no cover
        raise RuntimeError("Do odczytu zgód wymagany jest openpyxl.") from exc

    source = Path(path)
    if not source.exists():
        raise FileNotFoundError(source)
    workbook = load_workbook(source, read_only=True, data_only=False)
    try:
        selected = None
        positions = None
        for sheet in workbook.worksheets:
            try:
                positions = _header_positions(sheet.iter_rows(values_only=True))
            except ValueError:
                continue
            selected = sheet
            break
        if selected is None or positions is None:
            raise ValueError(
                "Żaden arkusz pliku zgód nie zawiera wymaganych nagłówków."
            )
        header_row, code_column, branch_column = positions
        by_branch: Dict[str, set[str]] = {key: set() for key in BRANCH_ORDER}
        owner: Dict[str, str] = {}
        source_rows = 0
        duplicates = 0
        maximum_column = max(code_column, branch_column) + 1
        for row_number, row in enumerate(
            selected.iter_rows(
                min_row=header_row + 1,
                max_col=maximum_column,
                values_only=True,
            ),
            start=header_row + 1,
        ):
            raw_code = row[code_column]
            raw_branch = row[branch_column]
            if raw_code in (None, "") and raw_branch in (None, ""):
                continue
            source_rows += 1
            if not isinstance(raw_code, str) or not raw_code.strip() or raw_code.startswith("="):
                raise ValueError(
                    f"{selected.title}!{row_number}: Kod PP Aktualny musi być "
                    "niepustym tekstem (nie liczbą Excela)."
                )
            code = raw_code.strip()
            branch = canonical_branch(raw_branch)
            previous = owner.get(code)
            if previous is not None:
                if previous != branch:
                    raise ValueError(
                        f"PPE {code!r} występuje w pliku zgód przy dwóch OSD."
                    )
                duplicates += 1
                continue
            owner[code] = branch
            by_branch[branch].add(code)
    finally:
        workbook.close()
    if not owner:
        raise ValueError("Plik zgód nie zawiera żadnego PPE.")
    return ConsentIndex(
        by_branch={key: frozenset(values) for key, values in by_branch.items()},
        source_rows=source_rows,
        duplicate_rows=duplicates,
    )


def load_consent_workbook(
    path: str | Path, osd_mapping_path: str | Path | None = None,
) -> ConsentIndex:
    """Aktualna lista wyznacza zgodę; stary eksport służy tylko do mapowania OSD.

    W nowym układzie niepusta kolumna korekty ma bezwzględne pierwszeństwo.
    Błędna korekta zatrzymuje import, bez powrotu do nieaktualnego kodu A.
    Identyfikatory muszą być tekstem, aby zachować wszystkie cyfry i zera.
    """
    from openpyxl import load_workbook

    workbook = load_workbook(Path(path), read_only=True, data_only=False)
    try:
        selected = None
        positions = None
        for sheet in workbook.worksheets:
            for number, row in enumerate(sheet.iter_rows(values_only=True), 1):
                names = [_plain(value) for value in row]
                if "numer ppe" in names:
                    corrected = next((i for i, value in enumerate(names) if value ==
                                      "wlasciwy ppe w przypadku bledow w kolumnie a"), None)
                    if corrected is None:
                        raise ValueError("Nowy plik zgód musi zawierać kolumnę 'Właściwy PPE (w przypadku błędów w kolumnie A)'.")
                    selected, positions = sheet, (number, names.index("numer ppe"), corrected)
                    break
                if number >= 30:
                    break
            if selected is not None:
                break
        if selected is None:
            return _load_legacy_consent_workbook(path)
        if osd_mapping_path is None:
            raise ValueError("Nowa lista PPE nie zawiera OSD. Wskaż pierwotny plik zgód przez --osd-mapping lub w oknie programu.")
        mapping = _load_legacy_consent_workbook(osd_mapping_path)
        owners = {code: branch for branch, codes in mapping.by_branch.items() for code in codes}
        header, original_column, corrected_column = positions
        by_branch = {key: set() for key in BRANCH_ORDER}
        seen = set()
        missing_rows = []
        source_rows = duplicates = corrected_rows = 0
        for number, row in enumerate(selected.iter_rows(
            min_row=header + 1, max_col=max(original_column, corrected_column) + 1,
            values_only=True,
        ), header + 1):
            original, corrected = row[original_column], row[corrected_column]
            if not _has_value(original) and not _has_value(corrected):
                continue
            source_rows += 1
            use_corrected = _has_value(corrected)
            raw_code = corrected if use_corrected else original
            corrected_rows += int(use_corrected)
            if not isinstance(raw_code, str) or raw_code.startswith("="):
                raise ValueError(f"{selected.title}!{number}: używany kod PPE musi być tekstem, nie liczbą ani formułą Excela.")
            code = raw_code.strip()
            if code in seen:
                duplicates += 1
                continue
            seen.add(code)
            if code not in owners:
                missing_rows.append(number)
            else:
                by_branch[owners[code]].add(code)
        if missing_rows:
            raise ValueError(f"Brak przypisania OSD w pierwotnym pliku dla {len(missing_rows)} PPE. "
                             f"Wiersze nowej listy (pierwsze 10): {missing_rows[:10]}. Popraw mapowanie przed obliczeniami.")
        if not seen:
            raise ValueError("Plik zgód nie zawiera żadnego PPE.")
        return ConsentIndex({key: frozenset(values) for key, values in by_branch.items()},
                            source_rows, duplicates, corrected_rows)
    finally:
        workbook.close()
